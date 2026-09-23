import hx
import os
import re
import html as _html
from html.parser import HTMLParser
from io import BytesIO
from copy import deepcopy

from mailmerge import MailMerge

from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _extract_data_uri(src: str):
    s = (src or "").strip()
    if not s:
        return None
    m = re.match(r"data:image/(png|jpeg|jpg|gif);base64,(.*)$", s, re.IGNORECASE | re.DOTALL)
    if not m:
        return None
    b64 = m.group(2)
    try:
        import base64
        return base64.b64decode(b64)
    except Exception:
        return None


class _HTMLToBlocks(HTMLParser):
    """
    Blocks:
      {"type": "p", "runs": [{"text":..,"bold":..,"italic":..,"underline":..,"size_px":..}]}
      {"type": "list", "ordered": bool, "items": [{"runs": [...], "level": int}, ...]}
      {"type": "table", "rows": [[cellRuns, cellRuns, ...], ...]}
      {"type": "img", "bytes": b"...", "width_in": float}
    """
    def __init__(self):
        super().__init__()
        self.blocks = []
        self._stack = []
        self._cur_runs = None
        self._font_size = None
        self._bold = False
        self._italic = False
        self._underline = False

        self._list_stack = []
        self._cur_list = None
        self._cur_item_runs = None
        self._cur_item_level = 0

        self._cur_table = None
        self._cur_row = None

    def handle_starttag(self, tag, attrs):
        tag = (tag or "").lower()
        attrs = dict(attrs or [])
        self._stack.append(tag)

        if tag in ("p", "div", "h2", "h3", "h4"):
            self._start_paragraph()
        elif tag == "br":
            if any(t in ("li", "td", "th") for t in self._stack):
                self._emit_text(" ")
            else:
                self._emit_text("\n")
        elif tag in ("strong", "b"):
            self._bold = True
        elif tag in ("em", "i"):
            self._italic = True
        elif tag == "u":
            self._underline = True
        elif tag in ("ul", "ol"):
            self._flush_p_if_any()
            self._start_list(ordered=(tag == "ol"))
        elif tag == "li":
            self._start_list_item()
        elif tag == "span":
            st = attrs.get("style", "")
            if st:
                m = re.search(r"font-size\s*:\s*([0-9\.]+)px", st, re.I)
                if m:
                    try:
                        self._font_size = float(m.group(1))
                    except Exception:
                        pass
        elif tag == "table":
            self._flush_p_if_any()
            self._cur_table = []
        elif tag == "tr":
            if self._cur_table is not None:
                self._cur_row = []
        elif tag in ("td", "th", "tbody"):
            self._start_paragraph()
        elif tag == "img":
            src = attrs.get("src", "") or ""
            width_in = None
            st = attrs.get("style", "")
            if st:
                m = re.search(r"width\s*:\s*([0-9\.]+)px", st, re.I)
                if m:
                    try:
                        width_in = float(m.group(1)) / 96.0
                    except Exception:
                        pass
            if width_in is None:
                w_attr = attrs.get("width")
                if w_attr:
                    try:
                        width_in = float(re.sub(r"[^\d.]", "", str(w_attr))) / 96.0
                    except Exception:
                        pass
            if width_in is None:
                width_in = 5.5

            data = _extract_data_uri(src)
            if data:
                self._flush_p_if_any()
                self.blocks.append({"type": "img", "bytes": data, "width_in": width_in})

    def handle_endtag(self, tag):
        tag = (tag or "").lower()

        if tag in ("p", "div", "h2", "h3", "h4"):
            self._finish_paragraph()
        elif tag in ("strong", "b"):
            self._bold = False
        elif tag in ("em", "i"):
            self._italic = False
        elif tag == "u":
            self._underline = False
        elif tag == "span":
            self._font_size = None
        elif tag == "li":
            self._finish_list_item()
        elif tag in ("ul", "ol"):
            self._finish_list()
        elif tag in ("td", "th"):
            self._finish_paragraph(cell=True)
        elif tag == "tr":
            if self._cur_table is not None and self._cur_row is not None:
                self._cur_table.append(self._cur_row)
                self._cur_row = None
        elif tag == "table":
            if self._cur_table is not None:
                self.blocks.append({"type": "table", "rows": self._cur_table})
                self._cur_table = None

        if self._stack and self._stack[-1] == tag:
            self._stack.pop()

    def handle_data(self, data):
        if not data:
            return
        self._emit_text(_html.unescape(data))

    def _emit_text(self, text):
        if self._cur_runs is None:
            self._start_paragraph()

        if text is None:
            return

        text = text.replace("\u00A0", " ")


        if not self._cur_runs:
            text = text.lstrip(" \t\r\n")

        if text == "":
            return

        self._cur_runs.append({
            "text": text,
            "bold": self._bold,
            "italic": self._italic,
            "underline": self._underline,
            "size_px": self._font_size
        })

    def _start_paragraph(self):
        if self._cur_runs is None:
            self._cur_runs = []

    def _flush_p_if_any(self):
        if self._cur_runs is not None:
            self.blocks.append({"type": "p", "runs": self._cur_runs})
            self._cur_runs = None

    def _finish_paragraph(self, cell=False):
        if self._cur_runs is None:
            if cell and self._cur_row is not None:
                self._cur_row.append([{"text": ""}])
            return
        runs = self._cur_runs
        self._cur_runs = None
        if cell and self._cur_row is not None:
            self._cur_row.append(runs)
        else:
            self.blocks.append({"type": "p", "runs": runs})

    def _start_list(self, ordered=False):
        self._list_stack.append('ol' if ordered else 'ul')
        if self._cur_list is None:
            self._cur_list = {"type": 'ol' if ordered else 'ul', "items": []}

    def _finish_list(self):
        if self._list_stack:
            self._list_stack.pop()
        if not self._list_stack and self._cur_list is not None:
            ordered = (self._cur_list["type"] == "ol")
            self.blocks.append({"type": "list", "ordered": ordered, "items": self._cur_list["items"]})
            self._cur_list = None

    def _start_list_item(self):
        self._cur_item_runs = []
        self._cur_runs = self._cur_item_runs
        self._cur_item_level = max(0, sum(1 for t in self._stack if t in ("ul", "ol")) - 1)

    def _finish_list_item(self):
        if self._cur_list is not None:
            runs = self._cur_item_runs or []
            self._cur_list["items"].append({"runs": runs, "level": self._cur_item_level})
        self._cur_item_runs = None
        self._cur_runs = None


def _add_zero_spacing_pprops(p_el):
    """
    Normalizes paragraph formatting for injected content:
    - removes first-line/left/right indentation (fixes "first paragraph indented")
    - adds a small paragraph spacing after (fixes "paragraphs glued")
    """
    ppr = p_el.find(qn('w:pPr'))
    if ppr is None:
        ppr = OxmlElement('w:pPr')
        p_el.insert(0, ppr)

    jc = ppr.find(qn('w:jc'))
    if jc is None:
        jc = OxmlElement('w:jc')
        ppr.append(jc)
    jc.set(qn('w:val'), 'left')

    ind = ppr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        ppr.append(ind)
    ind.set(qn('w:left'), '0')
    ind.set(qn('w:right'), '0')
    ind.set(qn('w:firstLine'), '0')
    ind.set(qn('w:hanging'), '0')

    sp = ppr.find(qn('w:spacing'))
    if sp is None:
        sp = OxmlElement('w:spacing')
        ppr.append(sp)
    sp.set(qn('w:before'), '0')
    sp.set(qn('w:after'), '160')  
    sp.set(qn('w:beforeAutospacing'), '0')
    sp.set(qn('w:afterAutospacing'), '0')

    return ppr


def _make_p_element_from_runs(runs):
    p = OxmlElement('w:p')
    _add_zero_spacing_pprops(p)
    for r in runs:
        text = r.get("text", "")
        if text == "":
            continue

        run = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')

        if r.get("bold"):
            rpr.append(OxmlElement('w:b'))
        if r.get("italic"):
            rpr.append(OxmlElement('w:i'))
        if r.get("underline"):
            u = OxmlElement('w:u')
            u.set(qn('w:val'), 'single')
            rpr.append(u)

        if r.get("size_hpt") is not None:
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), str(int(r["size_hpt"])))
            rpr.append(sz)
        else:
            sz_px = r.get("size_px")
            if sz_px:
                pt_half = int(round((sz_px * 72.0 / 96.0) * 2))
                sz = OxmlElement('w:sz')
                sz.set(qn('w:val'), str(pt_half))
                rpr.append(sz)

        color = r.get("color")
        if color:
            col_el = OxmlElement('w:color')
            col_el.set(qn('w:val'), color.lstrip('#'))
            rpr.append(col_el)

        if len(rpr):
            run.append(rpr)

        t = OxmlElement('w:t')
        if re.search(r'\s', text):
            t.set(qn('xml:space'), 'preserve')
        t.text = text
        run.append(t)
        p.append(run)

    if len(p) == 0:
        p.append(OxmlElement('w:r'))
    return p


def _apply_keep_with_next(p_el, keep_together=True):
    ppr = _add_zero_spacing_pprops(p_el)
    kw = OxmlElement('w:keepNext')
    ppr.append(kw)
    if keep_together:
        kt = OxmlElement('w:keepLines')
        ppr.append(kt)


def _apply_keep_chain(elements):
    for el in elements[:-1]:
        if el.tag == qn('w:p'):
            _apply_keep_with_next(el)


def _insert_table_at(docx_doc, parent, idx, table_rows):
    rows = len(table_rows)
    cols = max((len(r) for r in table_rows), default=0)
    if rows == 0 or cols == 0:
        parent.insert(idx, _make_p_element_from_runs([{"text": ""}]))
        return

    tbl = docx_doc.add_table(rows=rows, cols=cols)
    try:
        tbl.style = 'Table Grid'
    except Exception:
        pass

    for r_i, row in enumerate(table_rows):
        for c_i in range(cols):
            cell = tbl.cell(r_i, c_i)
            for p in list(cell.paragraphs):
                p._element.getparent().remove(p._element)
            runs = row[c_i] if c_i < len(row) else [{"text": ""}]
            cell._tc.append(_make_p_element_from_runs(runs))

    body = docx_doc._element.body
    t_el = tbl._tbl
    body.remove(t_el)
    parent.insert(idx, t_el)


def _insert_image_at(docx_doc, parent, idx, img_bytes, width_in):
    tmp_p = docx_doc.add_paragraph()
    tmp_run = tmp_p.add_run()
    try:
        tmp_run.add_picture(BytesIO(img_bytes), width=Inches(float(width_in)))
    except Exception:
        new_p = OxmlElement('w:p')
        _add_zero_spacing_pprops(new_p)
        parent.insert(idx, new_p)
        body = docx_doc._element.body
        body.remove(tmp_p._p)
        return

    new_p = OxmlElement('w:p')
    _add_zero_spacing_pprops(new_p)
    new_run = deepcopy(tmp_run._element)
    new_p.append(new_run)

    body = docx_doc._element.body
    body.remove(tmp_p._p)
    parent.insert(idx, new_p)


def _discover_numbering_seeds(docx_doc):
    seeds = {'ul': None, 'ol': None}
    for p in docx_doc.paragraphs:
        pPr = getattr(p._p, 'pPr', None)
        if pPr is None:
            continue
        numPr = pPr.find(qn('w:numPr'))
        if numPr is None:
            continue
        numId_el = numPr.find(qn('w:numId'))
        if numId_el is None:
            continue
        num_id = numId_el.get(qn('w:val'))
        style = (p.style.name or '').lower() if p.style else ''
        text = p.text or ''
        if seeds['ul'] is None and ('bullet' in style or '•' in text):
            seeds['ul'] = num_id
        elif seeds['ol'] is None and ('number' in style or re.match(r'^\s*\d+[\.\)]\s*', text)):
            seeds['ol'] = num_id
        if seeds['ul'] and seeds['ol']:
            break
    if not seeds['ul'] and seeds['ol']:
        seeds['ul'] = seeds['ol']
    if not seeds['ol'] and seeds['ul']:
        seeds['ol'] = seeds['ul']
    return seeds


def _make_list_p_from_runs(runs, num_id, ilvl):
    p = _make_p_element_from_runs(runs)
    ppr = _add_zero_spacing_pprops(p)

    numPr = ppr.find(qn('w:numPr'))
    if numPr is None:
        numPr = OxmlElement('w:numPr')
        ppr.append(numPr)

    ilvl_el = numPr.find(qn('w:ilvl'))
    if ilvl_el is None:
        ilvl_el = OxmlElement('w:ilvl')
        numPr.append(ilvl_el)
    ilvl_el.set(qn('w:val'), str(max(0, int(ilvl))))

    numId_el = numPr.find(qn('w:numId'))
    if numId_el is None:
        numId_el = OxmlElement('w:numId')
        numPr.append(numId_el)
    numId_el.set(qn('w:val'), str(num_id))

    return p


def _replace_token_anywhere(doc, token, html_value):
    NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

    parser = _HTMLToBlocks()
    parser.feed(html_value if isinstance(html_value, str) else ("" if html_value is None else str(html_value)))
    parser._flush_p_if_any()
    if parser._cur_list is not None or parser._list_stack:
        parser._finish_list()
    blocks = parser.blocks or [{"type": "p", "runs": [{"text": ""}]}]

    body_el = doc.part.element
    seeds = _discover_numbering_seeds(doc)

    for t in body_el.iter():
        if t.tag != qn('w:t'):
            continue
        if token not in (t.text or ""):
            continue
        p_anc = t
        while p_anc is not None and p_anc.tag != qn('w:p'):
            p_anc = p_anc.getparent()
        if p_anc is None:
            continue
        return _inject_blocks_at_paragraph(doc, p_anc, blocks, seeds)

    for p in body_el.findall('.//w:p', namespaces=NS):
        all_t = p.findall('.//w:t', namespaces=NS)
        if not all_t:
            continue
        joined = "".join((x.text or "") for x in all_t)
        if token not in joined:
            continue
        for x in all_t:
            x.text = ""
        return _inject_blocks_at_paragraph(doc, p, blocks, seeds)

    return False


def _inject_blocks_at_paragraph(doc, p_anc, blocks, seeds):
    parent = p_anc.getparent()
    if parent is None:
        return False

    target_parent = parent
    target_idx = parent.index(p_anc)
    parent.remove(p_anc)

    new_elements = []
    for b in blocks:
        if b["type"] == "p":
            new_elements.append(_make_p_element_from_runs(b["runs"]))
        elif b["type"] == "table":
            new_elements.append(OxmlElement('w:p'))
        elif b["type"] == "img":
            new_elements.append(OxmlElement('w:p'))
        elif b["type"] == "list":
            ordered = b.get("ordered", False)
            num_id = seeds['ol' if ordered else 'ul']
            if num_id:
                for it in b["items"]:
                    p_el = _make_list_p_from_runs(it["runs"], num_id=num_id, ilvl=it.get("level", 0))
                    new_elements.append(p_el)
            else:
                counter = 1
                for it in b["items"]:
                    runs = [{"text": (f"{counter}. " if ordered else "• ")}] + it["runs"]
                    p_el = _make_p_element_from_runs(runs)
                    new_elements.append(p_el)
                    if ordered:
                        counter += 1

    if new_elements and new_elements[0].tag == qn('w:p'):
        _apply_keep_with_next(new_elements[0])
    _apply_keep_chain(new_elements)

    idx = target_idx
    for el_new, b in zip(new_elements, blocks + [None] * max(0, len(new_elements) - len(blocks))):
        if b and b["type"] == "table":
            _insert_table_at(doc, target_parent, idx, b["rows"])
        elif b and b["type"] == "img":
            _insert_image_at(doc, target_parent, idx, b["bytes"], b.get("width_in", 5.5))
        else:
            target_parent.insert(idx, el_new)
        idx += 1

    return True


def generate_uw_rationale_doc(hxd, progress):
    template = os.path.join(os.path.dirname(__file__), "uw_rationale_doc.docx")
    mm = MailMerge(template)

    BD_TOKEN = "__BROKER_DISCUSSION_HTML__"
    UCOV_TOKEN = "__UNUSUAL_COVERAGE_HTML__"
    OF_TOKEN = "__OTHER_FACTORS_HTML__"
    ESG_TOKEN = "__ESG_HTML__"
    BPI_TOKEN = "__BPI_COMMENT_HTML__"

    mm.merge(
        insured_name=str(hxd.cds.standard_fields.insured_name or ""),
        inception_date=hxd.hx_core.inception_date.strftime("%d-%b-%Y") if hxd.hx_core.inception_date else "",
        expiry_date=hxd.hx_core.expiry_date.strftime("%d-%b-%Y") if hxd.hx_core.expiry_date else "",
        underwriter=str(hxd.cds.standard_fields.underwriter or ""),

        broker_discussion=BD_TOKEN,
        unusual_coverage=UCOV_TOKEN,
        other_factors=OF_TOKEN,
        esg=ESG_TOKEN,
        bpi_comment=BPI_TOKEN,
    )

    buf = BytesIO()
    mm.write(buf)
    buf.seek(0)

    docx_doc = Document(buf)

    broker_discussion_html = str(getattr(hxd.cds.uw_rationale, "broker_discussion", "") or "")
    unusual_coverage_html = str(getattr(hxd.cds.uw_rationale, "unusual_coverage", "") or "")
    other_factors_html = str(getattr(hxd.cds.uw_rationale, "other_factors", "") or "")
    esg_html = str(getattr(hxd.cds.uw_rationale, "esg", "") or "")
    bpi_comment_html = str(getattr(hxd.cds.uw_rationale, "bpi_comment", "") or "")

    _replace_token_anywhere(docx_doc, BD_TOKEN, broker_discussion_html)
    _replace_token_anywhere(docx_doc, UCOV_TOKEN, unusual_coverage_html)
    _replace_token_anywhere(docx_doc, OF_TOKEN, other_factors_html)
    _replace_token_anywhere(docx_doc, ESG_TOKEN, esg_html)
    _replace_token_anywhere(docx_doc, BPI_TOKEN, bpi_comment_html)

    with hxd.cds.uw_rationale.file_download_1.open("b") as f:
        docx_doc.save(f)