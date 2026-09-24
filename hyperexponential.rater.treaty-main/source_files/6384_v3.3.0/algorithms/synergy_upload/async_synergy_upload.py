import hx
import os
import msal
import json
import requests
import openpyxl
from openpyxl.utils import rows_from_range, cols_from_range
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.utils.cell import absolute_coordinate
from tempfile import NamedTemporaryFile
import base64
import datetime

def synergy_send_front_sheet(hxd, progress):

    # Load workbook
    workbook = load_workbook("front_sheet_template.xlsx")

    # Loop through named ranges and populate from hxd - named range name must match full hxd path
    write_to_named_ranges_from_hxd(hxd, workbook)

    # Set 'Bind Date' to today's date
    update_named_range_value(workbook, "non_hxd.bind_date", datetime.date.today())

    # Remove all layers where blank reference
    num_layers = len(hxd.cds.layers)
    max_layers = 8
    for i in range(1, max_layers + 1):
        reference = None
        if i <= num_layers:
            reference = hxd.cds.layers[i-1].section_reference
        if reference == None:
            named_range_name = f"behavior.layers.{i}"
            row_count = get_range_from_name(workbook, named_range_name).size['rows']
            clear_named_range(workbook, named_range_name) # Remove layer
            if i < max_layers:
                for i2 in range(i + 1, max_layers + 1):
                    named_range_name2 = f"behavior.layers.{i2}"
                    move_named_range(workbook, named_range_name2, -row_count, 0) # Shift higher layers up
        
    # Remove all named ranges
    delete_all_named_ranges(workbook)

    # Save workbook to temp file
    tmp_file = NamedTemporaryFile()
    workbook.save(tmp_file.name)

    # Email recipients
    is_dev = (hx.secrets.environment_name == 'beazley-dev' or hx.secrets.environment_name == 'beazley-tst')
    recipient_email = "SlipDataSynergy_UAT@beazley.com" if is_dev else "SlipDataIntoSynergy@beazley.com"
    #TODO comment out below
    # recipient_email = hx.meta.user.email # Override with current user's email temporarily
    cc_recipient_email = hx.meta.user.email

    # Send email
    subject = "Synergy Upload - " + (hxd.cds.standard_fields.insured_name or "") + " - " + hxd.cds.standard_fields.inception_date.strftime("%d.%m.%Y")
    message_body = ""
    attachment_filename = subject + ".xlsx"

    if hxd.synergy_upload.to_user_only:
        send_email(cc_recipient_email, subject, message_body, cc_recipient_email, tmp_file.name, attachment_filename)
    else:
        send_email(recipient_email, subject, message_body, cc_recipient_email, tmp_file.name, attachment_filename)


def synergy_send_rate_change(hxd, progress):

    cds = hxd.cds

    # Load workbook
    if cds.multi_year == "Yes":
        workbook = load_workbook("synergy_supplementary_template_my.xlsx")
    else:
        workbook = load_workbook("synergy_supplementary_template.xlsx")

    # Loop through named ranges and populate from hxd - named range name must match full hxd path
    write_to_named_ranges_from_hxd(hxd, workbook)

    # Get policy reference
    reference = None
    for l in hxd.cds.layers:
        if l.section_reference is not None:
            reference = l.section_reference
    if reference is None:
        reference = "NOREF"

    # Save workbook to temp file
    tmp_file = NamedTemporaryFile()
    workbook.save(tmp_file.name)

    # Email recipients
    recipient_email = hxd.synergy_upload.rate_change.email_recipients
    if recipient_email is None:
        hx.errors.fatal("Please add an email recipient")
    cc_recipient_email = hx.meta.user.email

    # Send email
    subject = "Synergy Supplementary - " + reference + " - " + (hxd.cds.standard_fields.insured_name or "") + " - " + hxd.cds.standard_fields.inception_date.strftime("%d.%m.%Y")
    message_body = ""
    attachment_filename = subject + ".xlsx"
    send_email(recipient_email, subject, message_body, cc_recipient_email, tmp_file.name, attachment_filename)


def send_eso(hxd, progress):

    cds = hxd.cds

    # Load workbook
    workbook = load_workbook("eso_template.xlsx")

    # Loop through named ranges and populate from hxd - named range name must match full hxd path
    write_to_named_ranges_from_hxd(hxd, workbook)

    # Get policy reference
    reference = None
    for l in hxd.cds.layers:
        if l.section_reference is not None:
            reference = l.section_reference
    if reference is None:
        reference = "NOREF"

    # Save workbook to temp file
    tmp_file = NamedTemporaryFile()
    workbook.save(tmp_file.name)

    # Email recipients
    recipient_email = cds.eso_template.email_recipients
    if recipient_email is None:
        hx.errors.fatal("Please add an email recipient")
    cc_recipient_email = hx.meta.user.email

    # Send email
    subject = "ESO - " + reference + " - " + (hxd.cds.standard_fields.insured_name or "") + " - " + (hxd.cds.programme or "") + " - " + hxd.cds.standard_fields.inception_date.strftime("%d.%m.%Y")
    message_body = ""
    attachment_filename = subject + ".xlsx"
    send_email(recipient_email, subject, message_body, cc_recipient_email, tmp_file.name, attachment_filename)


def load_workbook(wb_name):
    base_path = os.path.dirname(__file__)
    wb_path = os.path.join(base_path, wb_name)
    workbook = openpyxl.load_workbook(wb_path, keep_vba=False)
    return workbook

def write_to_named_ranges_from_hxd(hxd, workbook):
    names = list(workbook.defined_names)
    for name in names:
        if name.startswith("hxd."):
            defined_name = workbook.defined_names[name]
            for sheet_name, coord in defined_name.destinations:
                sheet = workbook[sheet_name]
                val = get_hxd_attr(hxd, name.replace('hxd.', ''))
                if val is not None and val != "":
                    sheet[coord].value = val


def update_named_range_value(workbook, named_range_name, val):
    named_range = get_named_range(workbook, named_range_name)
    worksheet = workbook[named_range['sheet_name']]
    cell_range = named_range['cell_range']
    worksheet[cell_range].value = val


def get_hxd_attr(obj, attrs):
    for attr in attrs.split('.'):
        if attr.isdigit():
            if int(attr) <= len(obj):
                obj = obj[int(attr)-1]
            else:
                return None
        else:
            obj = getattr(obj, attr)
    return obj

def get_named_range(workbook, named_range_name):
    defined_name = workbook.defined_names[named_range_name]
    for sheet_name, cell_range in defined_name.destinations:
        return {'sheet_name': sheet_name, 'cell_range': cell_range}

def get_range_from_name(workbook, named_range_name):
    named_range = get_named_range(workbook, named_range_name)
    cell_range = CellRange(named_range['cell_range'])
    return cell_range

def clear_named_range(workbook, named_range_name):
    named_range = get_named_range(workbook, named_range_name)
    worksheet = workbook[named_range['sheet_name']]
    for column in cols_from_range(named_range['cell_range']):
        for cell_coordinates in column:
            worksheet[cell_coordinates].value = None
            worksheet[cell_coordinates]._style = worksheet['A1']._style
    del workbook.defined_names[named_range_name]

def move_named_range(workbook, named_range_name, rows, cols):
    named_range = get_named_range(workbook, named_range_name)
    worksheet = workbook[named_range['sheet_name']]
    worksheet.move_range(named_range['cell_range'], rows=rows, cols=cols, translate=True)
    update_named_range_reference(workbook, named_range_name, offset_range_str(named_range['sheet_name'], named_range['cell_range'], rows, cols))

def update_named_range_reference(workbook, named_range_name, reference):
    defined_name = workbook.defined_names[named_range_name]
    defined_name.value = reference

def offset_range_str(sheet_name, cell_range, rows, cols):
    range_adjust = CellRange(cell_range)
    range_adjust.shift(col_shift=cols, row_shift=rows)
    return f"'{sheet_name}'!{absolute_coordinate(range_adjust.coord)}"

def delete_all_named_ranges(workbook):
    named_ranges = list(workbook.defined_names)
    for defined_name in named_ranges:
        del workbook.defined_names[defined_name]

def send_email(recipient_email, subject, message_body, cc_recipient_email = None, attachment = None, attachment_filename = None):
    recipient_email_list = [{'emailAddress': {'Address': address.strip()}} for address in recipient_email.split(';')]
    cc_recipient_email_list = [{'emailAddress': {'Address': address.strip()}} for address in cc_recipient_email.split(';')]
    client_id = hx.secrets.graph_clientid
    client_secret = hx.secrets.graph_clientsecret
    tenant_id = "9a50eba8-7568-447a-bcb9-27a0d464aa80"
    user_upn  =  "noreply_globalrating@beazley.com"

    # Microsoft Graph API endpoint to send an email as the application (service principal)
    graph_url = f"https://graph.microsoft.com/v1.0/users/{user_upn}/sendMail"

    # Create a confidential client application
    app = msal.ConfidentialClientApplication(
        client_id=client_id,
        authority=f"https://login.microsoftonline.com/{tenant_id}",
        client_credential=client_secret,
    )

    # Get an access token
    result = app.acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])
    access_token = result.get("access_token")
    
    # Add attachments
    attachments_list = []
    if attachment is not None and attachment_filename is not None:
            with open(attachment, 'rb') as f:
                attachment_content = f.read()
            content = {
                "@odata.type":"#microsoft.graph.fileAttachment",
                "name":attachment_filename,
                "contentBytes":base64.b64encode(attachment_content).decode('utf-8')
            }
            attachments_list.append(content)

    # Create the email message
    email_message = {
        "message": {
            "subject": subject,
            "body": {
                "contentType": "Text",
                "content": message_body,
            },
            "toRecipients": recipient_email_list,
            "ccRecipients": cc_recipient_email_list,
            "attachments": attachments_list
        },
        "saveToSentItems":"true"
    }

    # Convert the message to JSON format
    email_data = json.dumps(email_message)

    # Send the email
    response = requests.post(
        graph_url,
        headers={
            "Authorization": "Bearer " + access_token,
            "Content-Type": "application/json",
            'User-Agent': 'My User Agent 1.0',
        },
        data=email_data,
    )

    # Check the response status code
    if not response.status_code == 202:
        hx.errors.validation(f"Failed to send email. Status code: {response.status_code}")