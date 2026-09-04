import hx, json, openpyxl
import pandas as pd
pd.set_option('display.max_rows', None)
from copy import copy


# helper function to preserve cell formats
def write_preserve_style(sheet, row, col, value):
    cell            = sheet.cell(row=row, column=col)

    font            = copy(cell.font)
    fill            = copy(cell.fill)
    border          = copy(cell.border)
    alignment       = copy(cell.alignment)
    number_format   = copy(cell.number_format)
    protection      = copy(cell.protection)
    # style           = copy(cell.style) # style is not needed - it is deprecated - what is needed is to set the style in excel

    cell.value = value

    cell.font           = font
    cell.fill           = fill
    cell.border         = border
    cell.alignment      = alignment
    cell.number_format  = number_format
    cell.protection     = protection
    # cell.style          = style       # style is not needed - it is deprecated - what is needed is to set the style in excel





# Generate policy document in Excel
def tsk_policy_to_excel(hxd, progress):    
    from openpyxl.utils.cell import coordinate_to_tuple
    
    # Load string containing data
    data = json.loads(hxd.policy_doc.data_dict)

    # setting the right policy url string
    if   hx.secrets.environment_name == 'beazley-dev':      data["policy_url"] = data["dev_policy_url"]
    elif hx.secrets.environment_name == 'beazley-tst':      data["policy_url"] = data["tst_policy_url"]
    else:                                                   data["policy_url"] = data["prd_policy_url"]

    # Write the dictionary values to the Excel template
    template_path = f"./model/algorithms/example_document_template.xlsx"

    
    ############################
    ### standard code       ####
    ############################
    
    # Load the workbook 
    workbook = openpyxl.load_workbook(template_path)

    # Fill the named ranges with data
    for key, value in data.items():
        if key in workbook.defined_names:
            # Get the cell corresponding to the named range
            cells = workbook.defined_names[key].destinations
            for sheet_name, coord in cells:
                # Get the relevant worksheet
                sheet = workbook[sheet_name]
                # If value is a list, write values dynamically
                if isinstance(value, list):
                    coord = coord.replace("$", "")
                    row, col = coordinate_to_tuple(coord)
                    for df_row_index, df_row in enumerate(value, start=row):
                        for df_col_index, (cell_key, cell_value) in enumerate(df_row.items(), start=col):
                            # Write each value into the corresponding cell
                            write_preserve_style(sheet=sheet, row=df_row_index, col=df_col_index, value=cell_value)
                            # sheet.cell(row=df_row_index, column=df_col_index, value=cell_value)
                else:
                    # Assign single value to the cell
                    sheet[coord] = value





    ############################
    ### hiding extraneous rows/columns/sheets ####
    ############################


    # hide rows on Countries when not populated
    worksheet         = workbook['Countries']
    start_row         = 10
    end_row           = 109
    num_row           = len(data["tbl_countries"])
    start_hidden_row  = num_row     +     start_row
    if start_hidden_row < end_row:      
        worksheet.row_dimensions.group(  start = start_hidden_row,      end = end_row,        hidden= True)


    # hide rows on Exposure Details when not populated
    worksheet         = workbook['Exposure Details']
    start_row         = 26
    end_row           = 125
    num_row           = len(data["tbl_ihs_score"])
    start_hidden_row  = num_row     +     start_row
    if start_hidden_row < end_row:      
        worksheet.row_dimensions.group(  start = start_hidden_row,      end = end_row,        hidden= True)

    # hide Sheet Construction when not sslected
    worksheet           = workbook['Construction']
    hide_me             = data["construction_coverage"] != "Yes"
    if hide_me:
        worksheet.sheet_state = 'hidden'

    # Should hide on Risk Information TO and WL
    worksheet           = workbook['Risk Information']
    start_row = 25
    end_row = 27
    num_row = 1
    start_hidden_row  = num_row     +     start_row
    if start_hidden_row < end_row:      
        worksheet.row_dimensions.group(  start = start_hidden_row,      end = end_row,        hidden= True)


    # # hide using data grouping the relevant cells
    # if data['product'] == 'Political Risk':                             # data['product'] ... 'product' here is the name in our custom dictionary not the hxd
    #     risk_info_sheet.row_dimensions.group(start=36, end=50, hidden=True)
    #     exposure_sheet.column_dimensions.group(start='A', end='Q', hidden=True)
    #     rating_sum_sheet.row_dimensions.group(start=36, end=53, hidden=True)
    # else:
    #     risk_info_sheet.row_dimensions.group(start=36, end=50, hidden=False)
    #     exposure_sheet.column_dimensions.group(start='R', end='AC', hidden=True)
    #     rating_sum_sheet.row_dimensions.group(start=54, end=69, hidden=True)


    # Save the filled template to a new file
    with hxd.policy_doc.output_file.open("b") as f:
        workbook.save(f)

    # Store task data to allow comparison if anything changes in rating
    hxd.policy_doc.task_data_dict = hxd.policy_doc.data_dict

