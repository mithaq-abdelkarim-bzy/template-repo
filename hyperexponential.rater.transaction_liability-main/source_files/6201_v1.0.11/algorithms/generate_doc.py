import hx
import json
from datetime import datetime
import pandas as pd
import numpy as np
import math as math
from operator import itemgetter
from mailmerge import MailMerge
import os
import re
import html as _html
from html.parser import HTMLParser
from io import BytesIO
from copy import deepcopy
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def allow_policy_doc_download(hxd):
    layers = hxd.cds.layers
    doc = hxd.policy_doc
    if any(layer.limit in [None, 0] for layer in layers):
        doc.premium_check = "Limit must be filled out for all layers"
        doc.show_premium_check = True
    elif all(layer.quoted_premium in [None, 0] for layer in layers):
        doc.premium_check = (
            "Gross Quoted Premium should be greater than 0. "
            "Please input a valid number in Rating Summary."
        )
        doc.show_premium_check = True
    else:
        doc.show_generate_button = True


class _HTMLToBlocks(HTMLParser):
    def __init__(self):
        super().__init__()
        self.blocks = []
        self._stack = []
        self._cur_runs = None
        self._font_size = None
        self._bold = False
        self._italic = False
        self._underline = False

        self._list_stack = []
        self._cur_list = None
        self._cur_item_runs = None
        self._cur_item_level = 0

        self._cur_table = None
        self._cur_row = None

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        attrs = dict(attrs or [])
        self._stack.append(tag)

        if tag in ("p", "div", "h2", "h3", "h4"):
            self._start_paragraph()
        elif tag == "br":
            if any(t in ("li", "td", "th") for t in self._stack):
                self._emit_text(" ")
            else:
                self._emit_text("\n")
        elif tag in ("strong", "b"):
            self._bold = True
        elif tag in ("em", "i"):
            self._italic = True
        elif tag == "u":
            self._underline = True
        elif tag in ("ul", "ol"):
            self._flush_p_if_any()
            self._start_list(ordered=(tag == "ol"))
        elif tag == "li":
            self._start_list_item()
        elif tag == "span":
            st = attrs.get("style", "")
            if st:
                m = re.search(r"font-size\s*:\s*([0-9\.]+)px", st, re.I)
                if m:
                    try:
                        self._font_size = float(m.group(1))
                    except Exception:
                        pass
        elif tag == "table":
            self._flush_p_if_any()
            self._cur_table = []
        elif tag == "tr":
            if self._cur_table is not None:
                self._cur_row = []
        elif tag in ("td", "th","tbody"):
            self._start_paragraph()
        elif tag == "img":
            src = attrs.get("src", "") or ""
            width_in = None
            st = attrs.get("style", "")
            if st:
                m = re.search(r"width\s*:\s*([0-9\.]+)px", st, re.I)
                if m:
                    try:
                        width_in = float(m.group(1)) / 96.0
                    except Exception:
                        pass
            if width_in is None:
                w_attr = attrs.get("width")
                if w_attr:
                    try:
                        width_in = float(re.sub(r"[^\d.]", "", str(w_attr))) / 96.0
                    except Exception:
                        pass
            if width_in is None:
                width_in = 5.5
            data = _extract_data_uri(src)
            if data:
                self._flush_p_if_any()
                self.blocks.append({"type": "img", "bytes": data, "width_in": width_in})

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ("p", "div", "h2", "h3", "h4"):
            self._finish_paragraph()
        elif tag in ("strong", "b"):
            self._bold = False
        elif tag in ("em", "i"):
            self._italic = False
        elif tag == "u":
            self._underline = False
        elif tag == "span":
            self._font_size = None
        elif tag == "li":
            self._finish_list_item()
        elif tag in ("ul", "ol"):
            self._finish_list()
        elif tag in ("td", "th"):
            self._finish_paragraph(cell=True)
        elif tag == "tr":
            if self._cur_table is not None and self._cur_row is not None:
                self._cur_table.append(self._cur_row)
                self._cur_row = None
        elif tag == "table":
            if self._cur_table is not None:
                self.blocks.append({"type": "table", "rows": self._cur_table})
                self._cur_table = None
        if self._stack and self._stack[-1] == tag:
            self._stack.pop()

    def handle_data(self, data):
        if not data:
            return
        self._emit_text(_html.unescape(data))

    # paragraph helpers
    def _emit_text(self, text):
        if self._cur_runs is None:
            self._start_paragraph()
        self._cur_runs.append({
            "text": text,
            "bold": self._bold,
            "italic": self._italic,
            "underline": self._underline,
            "size_px": self._font_size
        })

    def _start_paragraph(self):
        if self._cur_runs is None:
            self._cur_runs = []

    def _flush_p_if_any(self):
        if self._cur_runs is not None:
            self.blocks.append({"type": "p", "runs": self._cur_runs})
            self._cur_runs = None

    def _finish_paragraph(self, cell=False):
        if self._cur_runs is None:
            if cell and self._cur_row is not None:
                self._cur_row.append([{"text": ""}])
            return
        runs = self._cur_runs
        self._cur_runs = None
        if cell and self._cur_row is not None:
            self._cur_row.append(runs)
        else:
            self.blocks.append({"type": "p", "runs": runs})

    # list helpers
    def _start_list(self, ordered=False):
        self._list_stack.append('ol' if ordered else 'ul')
        if self._cur_list is None:
            self._cur_list = {"type": 'ol' if ordered else 'ul', "items": []}

    def _finish_list(self):
        if self._list_stack:
            self._list_stack.pop()
        if not self._list_stack and self._cur_list is not None:
            ordered = (self._cur_list["type"] == "ol")
            self.blocks.append({"type": "list", "ordered": ordered, "items": self._cur_list["items"]})
            self._cur_list = None

    def _finish_list_item(self):
        if self._cur_list is not None:
            runs = self._cur_item_runs or []
            self._cur_list["items"].append({"runs": runs, "level": self._cur_item_level})
        self._cur_item_runs = None
        self._cur_runs = None

    def _start_list_item(self):
        self._cur_item_runs = []
        self._cur_runs = self._cur_item_runs
        self._cur_item_level = max(0, sum(1 for t in self._stack if t in ("ul", "ol")) - 1)


def _extract_data_uri(src: str):
    s = (src or "").strip()
    if not s:
        return None
    m = re.match(r"data:image/(png|jpeg|jpg|gif);base64,(.*)$", s, re.IGNORECASE | re.DOTALL)
    if not m:
        return None
    b64 = m.group(2)
    try:
        import base64
        return base64.b64decode(b64)
    except Exception:
        return None


def _add_zero_spacing_pprops(p_el):
    ppr = p_el.find(qn('w:pPr'))
    if ppr is None:
        ppr = OxmlElement('w:pPr')
        p_el.insert(0, ppr)
    sp = ppr.find(qn('w:spacing'))
    if sp is None:
        sp = OxmlElement('w:spacing')
        ppr.append(sp)
    sp.set(qn('w:before'), '0')
    sp.set(qn('w:after'), '0')
    return ppr


def _make_p_element_from_runs(runs):
    p = OxmlElement('w:p')
    _add_zero_spacing_pprops(p)
    for r in runs:
        text = r.get("text", "")
        if text == "":
            continue
        run = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        if r.get("bold"):
            rpr.append(OxmlElement('w:b'))
        if r.get("italic"):
            rpr.append(OxmlElement('w:i'))
        if r.get("underline"):
            u = OxmlElement('w:u')
            u.set(qn('w:val'), 'single')
            rpr.append(u)
        if r.get("size_hpt") is not None:
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), str(int(r["size_hpt"])))
            rpr.append(sz)
        else:
            sz_px = r.get("size_px")
            if sz_px:
                pt_half = int(round((sz_px * 72.0 / 96.0) * 2))
                sz = OxmlElement('w:sz')
                sz.set(qn('w:val'), str(pt_half))
                rpr.append(sz)
        color = r.get("color")
        if color:
            col_el = OxmlElement('w:color')
            col_el.set(qn('w:val'), color.lstrip('#'))
            rpr.append(col_el)
        if len(rpr):
            run.append(rpr)
        t = OxmlElement('w:t')
        if re.search(r'\s', text):
            t.set(qn('xml:space'), 'preserve')
        t.text = text
        run.append(t)
        p.append(run)
    if len(p) == 0:
        p.append(OxmlElement('w:r'))
    return p


def _apply_keep_with_next(p_el, keep_together=True):
    ppr = _add_zero_spacing_pprops(p_el)
    kw = OxmlElement('w:keepNext')
    ppr.append(kw)
    if keep_together:
        kt = OxmlElement('w:keepLines')
        ppr.append(kt)


def _insert_table_at(docx_doc, parent, idx, table_rows):
    rows = len(table_rows)
    cols = max((len(r) for r in table_rows), default=0)
    if rows == 0 or cols == 0:
        parent.insert(idx, _make_p_element_from_runs([{"text": ""}]))
        return
    tbl = docx_doc.add_table(rows=rows, cols=cols)
    try:
        tbl.style = 'Table Grid'
    except Exception:
        pass
    for r_i, row in enumerate(table_rows):
        for c_i in range(cols):
            cell = tbl.cell(r_i, c_i)
            for p in list(cell.paragraphs):
                p._element.getparent().remove(p._element)
            runs = row[c_i] if c_i < len(row) else [{"text": ""}]
            cell._tc.append(_make_p_element_from_runs(runs))
    body = docx_doc._element.body
    t_el = tbl._tbl
    body.remove(t_el)
    parent.insert(idx, t_el)


def _insert_image_at(docx_doc, parent, idx, img_bytes, width_in):
    tmp_p = docx_doc.add_paragraph()
    tmp_run = tmp_p.add_run()
    try:
        tmp_run.add_picture(BytesIO(img_bytes), width=Inches(float(width_in)))
    except Exception:
        new_p = OxmlElement('w:p')
        _add_zero_spacing_pprops(new_p)
        parent.insert(idx, new_p)
        body = docx_doc._element.body
        body.remove(tmp_p._p)
        return
    new_p = OxmlElement('w:p')
    _add_zero_spacing_pprops(new_p)
    new_run = deepcopy(tmp_run._element)
    new_p.append(new_run)
    body = docx_doc._element.body
    body.remove(tmp_p._p)
    parent.insert(idx, new_p)


def _last_descendant_p(el):
    last_p = None
    for node in el.iter():
        if node.tag == qn('w:p'):
            last_p = node
    return last_p


def _clear_page_break_before(p_el):
    ppr = p_el.find(qn('w:pPr'))
    if ppr is None:
        return
    for child in list(ppr):
        if child.tag == qn('w:pageBreakBefore'):
            ppr.remove(child)


def _runs_from_p(p_el):
    runs = []
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    for r in p_el.findall('.//w:r', namespaces=ns):
        texts = [t.text or "" for t in r.findall('.//w:t', namespaces=ns)]
        if not texts:
            continue
        txt = "".join(texts)
        if txt == "":
            continue
        rpr = r.find('w:rPr', namespaces=ns)
        is_b = rpr is not None and rpr.find('w:b', namespaces=ns) is not None
        is_i = rpr is not None and rpr.find('w:i', namespaces=ns) is not None
        u_el = rpr.find('w:u', namespaces=ns) if rpr is not None else None
        is_u = u_el is not None and (u_el.get(qn('w:val')) or 'single') != 'none'
        color_el = rpr.find('w:color', namespaces=ns) if rpr is not None else None
        color = color_el.get(qn('w:val')) if color_el is not None else None
        sz_el = rpr.find('w:sz', namespaces=ns) if rpr is not None else None
        size_hpt = int(sz_el.get(qn('w:val'))) if sz_el is not None and sz_el.get(qn('w:val')) else None
        runs.append({
            "text": txt,
            "bold": bool(is_b),
            "italic": bool(is_i),
            "underline": bool(is_u),
            "color": color,
            "size_hpt": size_hpt
        })
    return runs or [{"text": ""}]


def _steal_prev_heading_as_blocks(p_anc):
    el, tc, tr, tbl = p_anc, None, None, None
    while el is not None:
        if el.tag == qn('w:tc'): tc = el
        if el.tag == qn('w:tr'): tr = el
        if el.tag == qn('w:tbl'):
            tbl = el
            break
        el = el.getparent()

    heading_blocks = []

    if tbl is not None and tr is not None:
        prev_tr = tr.getprevious()
        if prev_tr is not None and prev_tr.tag == qn('w:tr'):
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            p_list = prev_tr.findall('.//w:p', namespaces=ns)
            for p_el in p_list:
                runs = _runs_from_p(p_el)
                heading_blocks.append({"type": "p", "runs": runs})
            tbl.remove(prev_tr)
            return heading_blocks

    prev = p_anc.getprevious()
    if prev is not None:
        if prev.tag == qn('w:p'):
            runs = _runs_from_p(prev)
            heading_blocks.append({"type": "p", "runs": runs})
            prev.getparent().remove(prev)
        elif prev.tag == qn('w:tbl'):
            cand = _last_descendant_p(prev)
            if cand is not None and cand.getparent() is not None:
                runs = _runs_from_p(cand)
                heading_blocks.append({"type": "p", "runs": runs})
                cand.getparent().remove(cand)
    return heading_blocks


def _remove_prev_heading_or_row(p_anc):
    return False


def _discover_numbering_seeds(docx_doc):
    seeds = {'ul': None, 'ol': None}
    for p in docx_doc.paragraphs:
        pPr = getattr(p._p, 'pPr', None)
        if pPr is None:
            continue
        numPr = pPr.find(qn('w:numPr'))
        if numPr is None:
            continue
        numId_el = numPr.find(qn('w:numId'))
        if numId_el is None:
            continue
        num_id = numId_el.get(qn('w:val'))
        style = (p.style.name or '').lower() if p.style else ''
        text = p.text or ''
        if seeds['ul'] is None and ('bullet' in style or '•' in text):
            seeds['ul'] = num_id
        elif seeds['ol'] is None and ('number' in style or re.match(r'^\s*\d+[\.\)]\s*', text)):
            seeds['ol'] = num_id
        if seeds['ul'] and seeds['ol']:
            break
    if not seeds['ul'] and seeds['ol']:
        seeds['ul'] = seeds['ol']
    if not seeds['ol'] and seeds['ul']:
        seeds['ol'] = seeds['ul']
    return seeds


def _make_list_p_from_runs(runs, num_id, ilvl):
    p = _make_p_element_from_runs(runs)
    ppr = _add_zero_spacing_pprops(p)
    numPr = ppr.find(qn('w:numPr'))
    if numPr is None:
        numPr = OxmlElement('w:numPr')
        ppr.append(numPr)
    ilvl_el = numPr.find(qn('w:ilvl'))
    if ilvl_el is None:
        ilvl_el = OxmlElement('w:ilvl')
        numPr.append(ilvl_el)
    ilvl_el.set(qn('w:val'), str(max(0, int(ilvl))))
    numId_el = numPr.find(qn('w:numId'))
    if numId_el is None:
        numId_el = OxmlElement('w:numId')
        numPr.append(numId_el)
    numId_el.set(qn('w:val'), str(num_id))
    return p


def _apply_keep_chain(elements):
    for el in elements[:-1]:
        if el.tag == qn('w:p'):
            _apply_keep_with_next(el)


def _replace_token_anywhere(doc, token, html):
    NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

    def _is_empty_paragraph(p_el):
        if p_el is None or p_el.tag != qn('w:p'):
            return False
        for t in p_el.findall('.//w:t', namespaces=NS):
            if (t.text or "").strip():
                return False
        return True

    def _is_empty_table(tbl_el):
        if tbl_el is None or tbl_el.tag != qn('w:tbl'):
            return False
        for t in tbl_el.findall('.//w:t', namespaces=NS):
            if (t.text or "").strip():
                return False
        if tbl_el.find('.//w:drawing', namespaces=NS) is not None:
            return False
        return True

    def _vacuum_between(body, start_el, stop_el):
        cur = start_el.getnext()
        while cur is not None and cur is not stop_el:
            nxt = cur.getnext()
            if cur.tag == qn('w:p') and _is_empty_paragraph(cur):
                body.remove(cur)
            elif cur.tag == qn('w:tbl') and _is_empty_table(cur):
                body.remove(cur)
            cur = nxt

    def _top_block_and_body(p_el):
        cur = p_el
        parent = cur.getparent()
        while parent is not None and parent.tag != qn('w:body'):
            cur = parent
            parent = parent.getparent()
        return cur, parent

    def _move_prev_uw_banner_before_topblock(top_block, body):
        if body is None:
            return
        scan = top_block.getprevious()
        while scan is not None:
            if scan.tag == qn('w:tbl'):
                txt = "".join((t.text or "") for t in scan.iter() if t.tag == qn('w:t'))
                if 'UW Rationale' in txt:
                    prev = scan.getprevious()
                    if _is_empty_paragraph(prev):
                        body.remove(prev)
                    idx_here = body.index(top_block)
                    body.remove(scan)
                    body.insert(idx_here, scan)
                    _vacuum_between(body, scan, top_block)
                    return
            scan = scan.getprevious()

    parser = _HTMLToBlocks()
    parser.feed(html if isinstance(html, str) else ("" if html is None else str(html)))
    parser._flush_p_if_any()
    if parser._cur_list is not None or parser._list_stack:
        parser._finish_list()
    blocks = parser.blocks or [{"type": "p", "runs": [{"text": ""}]}]

    body_el = doc.part.element
    seeds = _discover_numbering_seeds(doc)

    for t in body_el.iter():
        if t.tag != qn('w:t'):
            continue
        if token not in (t.text or ""):
            continue
        p_anc = t
        while p_anc is not None and p_anc.tag != qn('w:p'):
            p_anc = p_anc.getparent()
        if p_anc is None:
            continue
        return _inject_blocks_at_paragraph(doc, p_anc, blocks, token, seeds,
                                           _top_block_and_body, _move_prev_uw_banner_before_topblock)

    for p in body_el.findall('.//w:p', namespaces=NS):
        all_t = p.findall('.//w:t', namespaces=NS)
        if not all_t:
            continue
        joined = "".join((x.text or "") for x in all_t)
        if token not in joined:
            continue
        for x in all_t:
            x.text = ""
        return _inject_blocks_at_paragraph(doc, p, blocks, token, seeds,
                                           _top_block_and_body, _move_prev_uw_banner_before_topblock)

    return False


def _inject_blocks_at_paragraph(doc, p_anc, blocks, token, seeds,
                                top_block_and_body_fn, move_banner_fn):
    parent = p_anc.getparent()
    if parent is None:
        return False

    heading_blocks = _steal_prev_heading_as_blocks(p_anc)
    if heading_blocks:
        blocks = heading_blocks + blocks

    if token == "__TD_HTML__":
        top_block, body = top_block_and_body_fn(p_anc)
        move_banner_fn(top_block, body)

    prev_sib = p_anc.getprevious()

    def _is_empty_paragraph_local(p_el):
        if p_el is None or p_el.tag != qn('w:p'):
            return False
        for t in p_el.iter():
            if t.tag == qn('w:t') and (t.text or "").strip():
                return False
        return True
    if _is_empty_paragraph_local(prev_sib):
        parent.remove(prev_sib)

    target_parent = parent
    target_idx = parent.index(p_anc)
    parent.remove(p_anc)

    if token == "__TD_HTML__":
        top_block, body = top_block_and_body_fn(
            target_parent[target_idx-1] if target_idx > 0 else target_parent[0]
        )
        scan = top_block.getprevious()
        banner = None
        while scan is not None:
            if scan.tag in (qn('w:tbl'), qn('w:p')):
                txt = "".join((t.text or "") for t in scan.iter() if t.tag == qn('w:t'))
                if "uw rationale" in txt.lower():
                    banner = scan
                    break
            scan = scan.getprevious()
        if banner is not None and body is not None:
            cur = banner.getnext()
            while cur is not None and cur is not top_block:
                nxt = cur.getnext()
                if cur.tag == qn('w:p') and _is_empty_paragraph_local(cur):
                    body.remove(cur)
                elif cur.tag == qn('w:tbl'):
                    has_text = any((t.text or "").strip() for t in cur.iter() if t.tag == qn('w:t'))
                    has_drawing = any(el.tag == qn('w:drawing') for el in cur.iter())
                    if not has_text and not has_drawing:
                        body.remove(cur)
                cur = nxt
            _clear_page_break_before(top_block)
            target_parent = body
            target_idx = body.index(banner) + 1

    new_elements = []
    for b in blocks:
        if b["type"] == "p":
            p_new = _make_p_element_from_runs(b["runs"])
            new_elements.append(p_new)
        elif b["type"] == "table":
            new_elements.append(OxmlElement('w:p'))
        elif b["type"] == "img":
            new_elements.append(OxmlElement('w:p'))
        elif b["type"] == "list":
            ordered = b.get("ordered", False)
            num_id = seeds['ol' if ordered else 'ul']
            if num_id:
                for it in b["items"]:
                    p_el = _make_list_p_from_runs(it["runs"], num_id=num_id, ilvl=it.get("level", 0))
                    new_elements.append(p_el)
            else:
                counter = 1
                for it in b["items"]:
                    runs = [{"text": (f"{counter}. " if ordered else "• ")}] + it["runs"]
                    p_el = _make_p_element_from_runs(runs)
                    new_elements.append(p_el)
                    if ordered:
                        counter += 1

    if new_elements and new_elements[0].tag == qn('w:p'):
        _apply_keep_with_next(new_elements[0])
    _apply_keep_chain(new_elements)

    idx = target_idx
    for el_new, b in zip(new_elements, blocks + [None] * max(0, len(new_elements) - len(blocks))):
        if b and b["type"] == "table":
            _insert_table_at(doc, target_parent, idx, b["rows"])
        elif b and b["type"] == "img":
            _insert_image_at(doc, target_parent, idx, b["bytes"], b.get("width_in", 5.5))
        else:
            target_parent.insert(idx, el_new)
        idx += 1

    return True


def _html_to_plaintext(html_str: str) -> str:
    if html_str is None:
        return ""
    s = str(html_str)

    class _Mini(HTMLParser):
        def __init__(self):
            super().__init__()
            self.out = []
            self.ol = []
            self.in_li = False
            self.in_cell = False
            self.first_cell = True

        def _nl(self):
            if self.out and not str(self.out[-1]).endswith("\n"):
                self.out.append("\n")

        def handle_starttag(self, tag, attrs):
            tag = tag.lower()
            if tag in ("p", "div", "h2", "h3"):
                if not self.in_li:
                    self._nl()
            elif tag == "br":
                self.out.append(" " if (self.in_li or self.in_cell) else "\n")
            elif tag == "ul":
                self.ol.append(None)
            elif tag == "ol":
                self.ol.append(1)
            elif tag == "li":
                self.in_li = True
                self._nl()
                if self.ol:
                    if self.ol[-1] is None:
                        self.out.append("• ")
                    else:
                        self.out.append(f"{self.ol[-1]}. ")
                        self.ol[-1] += 1
            elif tag == "tr":
                self._nl()
                self.first_cell = True
            elif tag in ("td", "th"):
                if not self.first_cell:
                    self.out.append("\t")
                self.first_cell = False
                self.in_cell = True

        def handle_endtag(self, tag):
            tag = tag.lower()
            if tag in ("ul", "ol"):
                if self.ol:
                    self.ol.pop()
                    self._nl()
            elif tag == "li":
                self.in_li = False
                self._nl()
            elif tag == "tr":
                self._nl()
            elif tag in ("td", "th"):
                self.in_cell = False
            elif tag in ("p", "div", "h2", "h3"):
                self._nl()

        def handle_data(self, data):
            if not data:
                return
            txt = _html.unescape(data)
            txt = re.sub(r"[ \t\u00A0]+", " ", txt)
            self.out.append(txt)

        def get(self):
            t = "".join(self.out)
            t = re.sub(r"\n{3,}", "\n\n", t).strip()
            return t

    if not re.search(r"<[A-Za-z]", s):
        return s.replace("\r\n", "\n")
    p = _Mini()
    p.feed(s)
    return p.get()


def generate_doc(hxd, progress):
    cds = hxd.cds
    sf = cds.standard_fields
    eso = cds.eso
    rat = cds.rationale
    pol_info = cds.policy_info
    freq = cds.rating_factors.freq
    sev = cds.rating_factors.sev

    underwriter = sf.underwriter
    currency = cds.currencies.source_currency

    project_name = pol_info.project_name if pol_info.project_name is not None else ""
    target_name = pol_info.target_name if pol_info.target_name is not None else ""
    transaction_value = f"{pol_info.transaction_value:,.0f}" if pol_info.transaction_value is not None else ""
    total_limit = f"{pol_info.total_limit:,.0f}" if pol_info.total_limit is not None else ""
    total_limit_percentage = f"{pol_info.total_limit_percentage * 100:.2f} %" if pol_info.total_limit_percentage is not None else ""
    brokerage = f"{float(pol_info.brokerage * 100):.2f} %" if pol_info.brokerage is not None else ""
    policyholder = pol_info.policyholder if pol_info.policyholder is not None else ""
    buyer_name = pol_info.buyer_name if pol_info.buyer_name is not None else ""
    buyer_lawyer = pol_info.buyer_lawyer if pol_info.buyer_lawyer is not None else ""
    seller_name = pol_info.seller_name if pol_info.seller_name is not None else ""
    seller_lawyer = pol_info.seller_lawyer if pol_info.seller_lawyer is not None else ""
    uw_expenses_flag = "Yes" if pol_info.uw_expenses_flag else "No"
    uw_expenses = f"{pol_info.uw_expenses:,.0f}" if pol_info.uw_expenses is not None else ""
    distress_bus_flag = "Yes" if pol_info.distressed_business_flag else "No"

    freq_target_jurisdiction = freq.target.jurisdiction if freq.target.jurisdiction is not None else ""
    freq_target_industry = freq.target.industry if freq.target.industry is not None else ""
    freq_target_business_nature = freq.target.business_nature if freq.target.business_nature is not None else ""
    freq_buyer_jurisdiction = freq.buyer.jurisdiction if freq.buyer.jurisdiction is not None else ""
    freq_buyer_industry = freq.buyer.industry if freq.buyer.industry is not None else ""
    freq_buyer_business_nature = freq.buyer.business_nature if freq.buyer.business_nature is not None else ""
    freq_default_score = f"{float(freq.default_score * 100):.2f} %" if freq.default_score is not None else ""
    freq_adjustment = f"{float(freq.freq_adjustment * 100):.2f} %" if freq.freq_adjustment is not None else ""
    freq_adjusted_freq = f"{float(freq.freq_adjusted * 100):.2f} %" if freq.freq_adjusted is not None else ""

    sev_dd_coverage_flag = "Yes" if sev.due_diligence.coverage_flag else "No"
    sev_dd_severity = sev.due_diligence.severity_assessment if sev.due_diligence.severity_assessment is not None else ""
    sev_dd_adj = f"{float(sev.due_diligence.adjustment * 100):.2f} %" if sev.due_diligence.adjustment is not None else ""

    sev_dis_cov_flag = "Yes" if sev.disclosure.coverage_flag else "No"
    sev_disc_severity = sev.disclosure.severity_assessment if sev.disclosure.severity_assessment is not None else ""
    sev_dis_adj = f"{float(sev.disclosure.adjustment * 100):.2f} %" if sev.disclosure.adjustment is not None else ""

    sev_gen_war_cov_flag = "Yes" if sev.general_warranties.coverage_flag else "No"
    sev_gen_war_severity = sev.general_warranties.severity_assessment if sev.general_warranties.severity_assessment is not None else ""
    sev_gen_war_adj = f"{float(sev.general_warranties.adjustment * 100):.2f} %" if sev.general_warranties.adjustment is not None else ""
    sev_gen_war_term = f"{sev.general_warranties.term:,.2f}" if sev.general_warranties.term is not None else ""
    sev_gen_war_term_mod = f"{sev.general_warranties.term_modifier:,.2f}" if sev.general_warranties.term_modifier is not None else ""
    sev_gen_war_comment = _html_to_plaintext(sev.general_warranties.comment)

    tax_warr = getattr(sev, "tax_warranties", None)
    sev_tax_war_cov_flag = "Yes" if (tax_warr and tax_warr.coverage_flag) else "No"
    sev_tax_war_severity = tax_warr.severity_assessment if (tax_warr and tax_warr.severity_assessment is not None) else ""
    sev_tax_war_adj = f"{float(tax_warr.adjustment * 100):.2f} %" if (tax_warr and tax_warr.adjustment is not None) else ""
    sev_tax_war_term = f"{tax_warr.term:,.2f}" if (tax_warr and tax_warr.term is not None) else ""
    sev_tax_war_term_mod = f"{tax_warr.term_modifier:,.2f}" if (tax_warr and tax_warr.term_modifier is not None) else ""
    sev_tax_war_comment = _html_to_plaintext(tax_warr.comment if tax_warr else "")

    fun_top_up_flag = "Yes" if pol_info.fundamental_top_up_flag else "No"
    fundamental_base_rate = f"{float(cds.rating_factors.fun_top_up.base_rate * 100):.2f} %" if pol_info.fundamental_top_up_flag else "N/A"
    fundamental_adj = f"{float(cds.rating_factors.fun_top_up.adjustment * 100):.2f} %" if pol_info.fundamental_top_up_flag else "N/A"
    if not pol_info.fundamental_top_up_flag:
        fundamental_term = "N/A"
        fundamental_comment_text = "N/A"
    else:
        fundamental_term = f"{cds.rating_factors.fun_top_up.term:,.2f}" if cds.rating_factors.fun_top_up.term else ""
        fundamental_comment_text = _html_to_plaintext(cds.rating_factors.fun_top_up.comment)

    policy_reference = eso.section_references
    bind_date = eso.bind_date.strftime("%d %b %Y") if eso.bind_date else ""
    premium = f"{eso.premium_total.selected:,.0f}" if eso.premium_total.selected is not None else ""
    eso_limit = f"{eso.limit_total.selected:,.0f}" if eso.limit_total.selected is not None else ""
    term = f"{eso.term.selected}" + (f"\n{_html_to_plaintext(eso.additional_term_info)}" if eso.additional_term_info else "")
    over_lining = eso.over_lining
    cob = eso.cob
    authorising_comments_text = _html_to_plaintext(eso.authorising_comments)
    approving_uw = eso.approving_uw
    authorisation_date = eso.authorisation_date.strftime("%d %b %Y") if eso.authorisation_date else ""

    td_html = rat.transaction_description or ""
    ki_html = rat.knowledge_of_insured or ""
    rw_html = rat.reasons_for_writing_risk or ""
    uc_html = rat.complex_considerations or ""
    ex_html = rat.exclusions_and_carve_outs or ""

    # Layers
    layers_copy = []
    for layer in cds.layers:
        layers_copy.append({
            "section_reference": f"{layer.section_reference}" if layer.section_reference else "",
            "is_fun_top_up_coverage": "Yes" if layer.is_fun_top_up_coverage else "No",
            "is_primary_excess": f"{layer.is_primary_excess}" if layer.is_primary_excess else "",
            "option_name": f"{layer.option_name}" if layer.option_name else "",
            "status": f"{layer.status}" if layer.status else "",
            "written_line": f"{float(layer.written_line * 100):.2f} %" if layer.written_line else "",
            "limit": f"{layer.limit:,.0f}" if layer.limit else "",
            "limit_pct": f"{float(layer.limit_pct * 100):.2f} %" if layer.limit_pct else "",
            "excess": f"{layer.excess.selected:,.0f}" if layer.excess.selected else "",
            "excess_pct": f"{float(layer.excess_pct.selected * 100):.2f} %" if layer.excess_pct.selected else "",
            "indicated": "Yes" if layer.indicated else "No",
            "expected_loss_cost": f"{layer.expected_loss_cost:,.0f}" if layer.expected_loss_cost else "",
            "model_rol": f"{float(layer.model_rol * 100):.2f} %" if layer.model_rol else "",
            "quoted_premium": f"{layer.quoted_premium:,.0f}" if layer.quoted_premium else "",
            "model_premium": f"{layer.model_premium:,.0f}",
            "technical_premium": f"{layer.technical_premium:,.0f}",
            "benchmark_premium": f"{layer.benchmark_premium:,.0f}",
            "tpi": f"{float(layer.tpi * 100):.2f} %" if layer.tpi else "",
            "bpi": f"{float(layer.bpi * 100):.2f} %" if layer.bpi else "",
            "pflr": f"{float(layer.pflr * 100):.2f} %" if layer.pflr else "",
            "roc": f"{float(layer.roc * 100):.2f} %" if layer.roc else "",
            "uw_adj_impact": f"{float(layer.uw_adj_impact * 100):.2f} %",
        })

    TD_TOKEN = "__TD_HTML__"
    KI_TOKEN = "__KI_HTML__"
    RW_TOKEN = "__RW_HTML__"
    UC_TOKEN = "__UC_HTML__"
    EX_TOKEN = "__EX_HTML__"

    template = "uw_doc_template.docx"
    template_path = f"./model/algorithms/policy_doc_templates/{template}"
    mm = MailMerge(template_path)
    mm.merge(
        underwriter=underwriter,
        currency=currency,
        project_name=project_name,
        target_name=target_name,
        transaction_value=transaction_value,
        total_limit=total_limit,
        total_limit_percentage=total_limit_percentage,
        brokerage=brokerage,
        policyholder=policyholder,
        buyer_name=buyer_name,
        buyer_lawyer=buyer_lawyer,
        seller_name=seller_name,
        seller_lawyer=seller_lawyer,
        uw_expenses_flag=uw_expenses_flag,
        uw_expenses=uw_expenses,
        distress_bus_flag=distress_bus_flag,
        freq_target_jurisdiction=freq_target_jurisdiction,
        freq_target_industry=freq_target_industry,
        freq_target_business_nature=freq_target_business_nature,
        freq_buyer_jurisdiction=freq_buyer_jurisdiction,
        freq_buyer_industry=freq_buyer_industry,
        freq_buyer_business_nature=freq_buyer_business_nature,
        freq_default_score=freq_default_score,
        freq_adjustment=freq_adjustment,
        freq_adjusted_freq=freq_adjusted_freq,
        sev_dd_coverage_flag=sev_dd_coverage_flag,
        sev_dis_cov_flag=sev_dis_cov_flag,
        sev_gen_war_cov_flag=sev_gen_war_cov_flag,
        sev_tax_war_cov_flag=sev_tax_war_cov_flag,
        sev_dd_severity=sev_dd_severity,
        sev_disc_severity=sev_disc_severity,
        sev_gen_war_severity=sev_gen_war_severity,
        sev_tax_war_severity=sev_tax_war_severity,
        sev_dd_adj=sev_dd_adj,
        sev_dis_adj=sev_dis_adj,
        sev_gen_war_adj=sev_gen_war_adj,
        sev_tax_war_adj=sev_tax_war_adj,
        sev_gen_war_term=sev_gen_war_term,
        sev_tax_war_term=sev_tax_war_term,
        sev_gen_war_term_mod=sev_gen_war_term_mod,
        sev_tax_war_term_mod=sev_tax_war_term_mod,
        sev_dd_comment=_html_to_plaintext(sev.due_diligence.comment),
        sev_dis_comment=_html_to_plaintext(sev.disclosure.comment),
        sev_gen_war_comment=sev_gen_war_comment,
        sev_tax_war_comment=sev_tax_war_comment,
        sev_modified_score=f"{sev.modified_score:,.2f}" if sev.modified_score is not None else "",
        sev_div_flag="Yes" if sev.div_flag else "No",
        sev_volitility=sev.volatility if sev.volatility is not None else "",
        sev_multiple=f"{sev.multiple:,.2f}" if sev.multiple is not None else "",
        fun_top_up_flag=fun_top_up_flag,
        fundamental_base_rate=f"{float(cds.rating_factors.fun_top_up.base_rate * 100):.2f} %" if pol_info.fundamental_top_up_flag else "N/A",
        fundamental_adj=f"{float(cds.rating_factors.fun_top_up.adjustment * 100):.2f} %" if pol_info.fundamental_top_up_flag else "N/A",
        fundamental_term=f"{cds.rating_factors.fun_top_up.term:,.2f}" if pol_info.fundamental_top_up_flag and cds.rating_factors.fun_top_up.term else "N/A",
        fundamental_comment=_html_to_plaintext(cds.rating_factors.fun_top_up.comment) if pol_info.fundamental_top_up_flag else "N/A",
        policy_reference=policy_reference,
        bind_date=bind_date,
        premium=premium,
        eso_limit=eso_limit,
        term=term,
        over_lining=over_lining,
        cob=cob,
        authorising_comments=authorising_comments_text,
        approving_uw=approving_uw,
        authorisation_date=authorisation_date,
        transaction_description=TD_TOKEN,
        knowledge_of_insured=KI_TOKEN,
        reasons_for_writing_risk=RW_TOKEN,
        complex_considerations=UC_TOKEN,
        exclusions_and_carve_outs=EX_TOKEN,
    )
    mm.merge_rows("option_name", layers_copy)
    mm.merge_rows("quoted_premium", layers_copy)

    buf = BytesIO()
    mm.write(buf)
    buf.seek(0)

    docx_doc = Document(buf)

    _replace_token_anywhere(docx_doc, TD_TOKEN, td_html)
    _replace_token_anywhere(docx_doc, KI_TOKEN, ki_html)
    _replace_token_anywhere(docx_doc, RW_TOKEN, rw_html)
    _replace_token_anywhere(docx_doc, UC_TOKEN, uc_html)
    _replace_token_anywhere(docx_doc, EX_TOKEN, ex_html)

    with hxd.policy_doc.output_file_doc.open("b") as out:
        docx_doc.save(out)
