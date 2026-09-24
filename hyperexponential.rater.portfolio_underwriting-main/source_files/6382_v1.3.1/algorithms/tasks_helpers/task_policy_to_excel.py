import hx, json, openpyxl
import pandas as pd
pd.set_option('display.max_rows', None)
from copy import copy
from openpyxl.utils.cell import coordinate_to_tuple
from datetime import datetime

# helper function to preserve cell formats
def write_preserve_style(sheet, row, col, value):
    cell            = sheet.cell(row=row, column=col)

    font            = copy(cell.font)
    fill            = copy(cell.fill)
    border          = copy(cell.border)
    alignment       = copy(cell.alignment)
    number_format   = copy(cell.number_format)
    protection      = copy(cell.protection)
    # style           = copy(cell.style) # style is not needed - it is deprecated - what is needed is to set the style in excel

    cell.value = value

    cell.font           = font
    cell.fill           = fill
    cell.border         = border
    cell.alignment      = alignment
    cell.number_format  = number_format
    cell.protection     = protection
    # cell.style          = style       # style is not needed - it is deprecated - what is needed is to set the style in excel


def hide_unused_rows(wb, sheet_name: str, start_row: int, num_rows_to_show: int, total_block_rows: int = 50) -> None:
    ws = wb[sheet_name]  # or wb.worksheets[...] if you prefer
    end_row = start_row + total_block_rows - 1
    # clamp so it behaves sensibly if num_rows_to_show is weird
    num_rows_to_show = max(0, min(num_rows_to_show, total_block_rows))
    start_hidden_row = start_row + num_rows_to_show
    if start_hidden_row <= end_row:
        for r in range(start_hidden_row, end_row + 1):
            ws.row_dimensions[r].hidden = True


def hide_sheet(wb, sheet_name: str, very=False):
    ws = wb[sheet_name]
    ws.sheet_state = "veryHidden" if very else "hidden"




# Generate policy document in Excel
def task_policy_to_excel(hxd):    
    
    # Load string containing data
    data = json.loads(hxd.non_cds.excel_analysis.data_dict)

    # setting the right policy url string
    if   hx.secrets.environment_name == 'beazley-dev':      data["summary_hx_model_ref"] = data.get("dev_policy_url", "")
    elif hx.secrets.environment_name == 'beazley-tst':      data["summary_hx_model_ref"] = data.get("tst_policy_url", "")
    else:                                                   data["summary_hx_model_ref"] = data.get("prd_policy_url", "")

    # settign the export dats time
    data['summary_export_date'] = datetime.now().strftime("%a %d %b %Y %H:%M:%S")

    # Write the dictionary values to the Excel template
    template_path = f"./model/algorithms/templates/analysis_overview.xlsx"
    
    ############################
    ### standard code       ####
    ############################
    
    # Load the workbook 
    workbook = openpyxl.load_workbook(template_path)

    # Fill the named ranges with data
    for key, value in data.items():
        if key in workbook.defined_names:
            # Get the cell corresponding to the named range
            cells = workbook.defined_names[key].destinations
            for sheet_name, coord in cells:
                # Get the relevant worksheet
                sheet = workbook[sheet_name]
                # If value is a list, write values dynamically
                if isinstance(value, list):
                    coord = coord.replace("$", "")
                    row, col = coordinate_to_tuple(coord)
                    for df_row_index, df_row in enumerate(value, start=row):
                        for df_col_index, (cell_key, cell_value) in enumerate(df_row.items(), start=col):
                            # Write each value into the corresponding cell
                            write_preserve_style(sheet=sheet, row=df_row_index, col=df_col_index, value=cell_value)
                            # sheet.cell(row=df_row_index, column=df_col_index, value=cell_value)
                else:
                    # Assign single value to the cell
                    sheet[coord] = value


    ############################
    ### hiding extraneous rows/columns/sheets ####
    ############################

    is_bbt          = data.get("is_bbt",         "No" ) == "Yes"
    is_case_priced  = data.get("is_case_priced", "Yes") == "Yes"        # default to is_case_priced
    is_prem_avail   = data.get("is_prem_avail",  "No" ) == "Yes"        
    is_pc           = data.get("summary_pc",     "Not Applicable" ) != "Not Applicable"

    if is_bbt or is_case_priced or (not is_prem_avail):
        hide_sheet(      workbook, 'Own Experience Analysis')


    if is_bbt or is_case_priced:
        hide_sheet(      workbook, 'SummaryTables')
        hide_unused_rows(workbook, 'Summary', 12, 0, 1)
        hide_unused_rows(workbook, 'Summary', 16, 0, 1)
        hide_unused_rows(workbook, 'Summary', 17, 0, 1)
        hide_unused_rows(workbook, 'Summary', 20, 0, 6)
        hide_unused_rows(workbook, 'Summary', 26, 0, 1)

    if not is_case_priced:
        hide_unused_rows(workbook, 'Summary', 31, 0, 1)

    if not (is_bbt or is_case_priced):
        num_lobs = len(data.get('table_model_gn_ulr',{}))
        hide_unused_rows(workbook,'SummaryTables',   9, num_lobs, 50)
        hide_unused_rows(workbook,'SummaryTables',  67, num_lobs, 50)
        hide_unused_rows(workbook,'SummaryTables', 127, num_lobs, 50)
        hide_unused_rows(workbook,'SummaryTables', 184, num_lobs, 50)
        hide_unused_rows(workbook,'SummaryTables', 240, num_lobs, 50)
        hide_unused_rows(workbook,'SummaryTables', 298, num_lobs, 50)
        hide_unused_rows(workbook,'SummaryTables', 354, num_lobs, 50)
        hide_unused_rows(workbook,'SummaryTables', 410, num_lobs, 50)
        hide_unused_rows(workbook,'SummaryTables',   9, num_lobs, 50)

    if not is_pc:
        hide_unused_rows(workbook,'SummaryTables', 407, 0, 55)        # pb request 17 march 2026

    # Save the filled template to a new file
    with hxd.non_cds.excel_analysis.output_file.open("b") as f:
        workbook.save(f)

    # Store task data to allow comparison if anything changes in rating
    hxd.non_cds.excel_analysis.task_data_dict = hxd.non_cds.excel_analysis.data_dict

