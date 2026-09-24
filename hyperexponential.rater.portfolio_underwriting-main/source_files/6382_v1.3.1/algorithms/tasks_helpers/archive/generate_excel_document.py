import os
import json
import openpyxl
from copy import copy
import pandas as pd
from io import StringIO
import algorithms.rate_utilities as utils
from openpyxl.utils.cell import coordinate_to_tuple
from functools import reduce
from algorithms.tasks_helpers.generate_excel_document_helpers import (
    build_key_lob_str,
    build_trifocus_string,
    build_risk_codes_str,
    build_prem_average_limit_str,
    build_gn_prem,
    build_deduction_strings,
    generate_profit_commission_description,
    build_anti_selection_uncertainty_str
)


def generate_excel_document(hxd):
    data_dict = create_excel_data_map(hxd)

    json_data = json.dumps(data_dict)

    workbook = import_workbook()

    write_values_to_excel(data_dict, workbook)

    hide_empty_rows(workbook, 43, 93, data_dict["model_gn_ulr_table"])
    hide_empty_rows(workbook, 101, 151, data_dict["selected_ulr_table"])
    hide_empty_rows(workbook, 161, 211, data_dict["tech_adequacy_metrics_table"])

    with hxd.cds.rationale.excel_rationale_template.open("b") as f:
        workbook.save(f)


def import_workbook():
    template = os.path.join(
        os.path.dirname(__file__),
        "..",
        "templates",
        "excel_rationale_template.xlsx"
    )

    workbook = openpyxl.load_workbook(template)

    return workbook


def write_values_to_excel(data, workbook):
    for key, value in data.items():
        if key in workbook.defined_names:
            # Get the cell corresponding to the named range
            cells = workbook.defined_names[key].destinations
            for sheet_name, coord in cells:
                # Get the relevant worksheet
                sheet = workbook[sheet_name]
                # If value is a list, write values dynamically
                if isinstance(value, list):
                    coord = coord.replace("$", "")
                    row, col = coordinate_to_tuple(coord)
                    for df_row_index, df_row in enumerate(value, start=row):
                        for df_col_index, (cell_key, cell_value) in enumerate(df_row.items(), start=col):
                            # Write each value into the corresponding cell
                            write_preserve_style(sheet=sheet, row=df_row_index, col=df_col_index, value=cell_value)
                            # sheet.cell(row=df_row_index, column=df_col_index, value=cell_value)
                else:
                    # Assign single value to the cell
                    sheet[coord] = value


def import_rater(hxd):
    cds = hxd.cds
    rs = cds.rating_summary
    pam = rs.pricing_adequacy_metrics

    rater = {}

    final_comp_str = hxd.non_cds.risk_code_composition.final_composition_str
    rater["risk_composition_final"] = pd.read_csv(StringIO(final_comp_str), keep_default_na=False)

    deductions_df_str = hxd.non_cds.assumed_deductions.final_deductions_str
    rater["deductions_df"] = pd.read_csv(StringIO(deductions_df_str))

    rater["prem_limit_df"] = utils.pd_df_from_hx_list_v2(
        cds.prem_limit_profile.table)
    
    rater["projected_gn_ulr_df"] = utils.pd_df_from_hx_list_v2(
        rs.model_gn_ulr.projected_gn_ulr.table)

    rater["model_weights_df"] = utils.pd_df_from_hx_list_v2(
        rs.model_gn_ulr.model_weights.table)

    rater["cat_loadings_df"] = utils.pd_df_from_hx_list_v2(
        rs.cat_loadings.cat_allocation)

    rater["add_loadings_df"] = utils.pd_df_from_hx_list_v2(
        rs.additional_loadings.additional_pricing_loads)

    pre_adj_df_str = hxd.non_cds.rating_summary.pricing_adequacy_metrics.pricing_adequacy_pre_adj.table_str
    rater["pre_adj"] = pd.read_csv(StringIO(pre_adj_df_str))

    act_basis_df_str = hxd.non_cds.rating_summary.pricing_adequacy_metrics.pricing_adequacy_actuarial_basis.table_str
    rater["act_basis"] = pd.read_csv(StringIO(act_basis_df_str))

    rater["final_pricing"] = utils.pd_df_from_hx_list_v2(
        pam.pricing_adequacy_final_pricing.table)
    
    return rater


def create_excel_data_map(hxd):

    rater = import_rater(hxd)

    key_risk_details = generate_key_risk_details_dict(hxd, rater)
    pricing_outputs = generate_pricing_outputs_dict(hxd, rater)
    model_gn_ulr_summary = generate_model_gn_ulr_summary(hxd)
    selected_ulr_summary = generate_selected_ulr_summary(hxd)
    tech_adequacy_metrics_summary = generate_tech_adequacy_metrics_summary(hxd)

    data_dict = (
        key_risk_details
        | pricing_outputs
        | model_gn_ulr_summary
        | selected_ulr_summary
        | tech_adequacy_metrics_summary
    )

    data_dict["model_gn_ulr_table"] = generate_model_gn_ulr_list_dict(hxd, rater)
    data_dict["selected_ulr_table"] = generate_selected_ulr_list_dict(hxd, rater)
    data_dict["tech_adequacy_metrics_table"] = generate_tech_adequacy_metrics_list_dict(hxd, rater)
    
    return data_dict
    

def generate_key_risk_details_dict(hxd, rater):
    cds = hxd.cds
    ri = cds.risk_information
    sf = cds.standard_fields

    key_lob_str = build_key_lob_str(rater["prem_limit_df"])
    trifocus_str = build_trifocus_string(rater["prem_limit_df"])
    risk_codes_str = build_risk_codes_str(rater["risk_composition_final"])
    renewal_str = "Renewal" if sf.is_renewal else "New Business"
    inception_date = sf.inception_date.strftime("%m/%d/%Y")
    expiry_date = sf.expiry_date.strftime("%m/%d/%Y")
    line_size = cds.prem_limit_profile.summary.bst_share_line_size
    prem_average_limit = build_prem_average_limit_str(hxd)
    gn_prem = build_gn_prem(hxd)
    deductions = build_deduction_strings(hxd, rater)
    pc_str = generate_profit_commission_description(hxd)

    data = {
        "insured_name": sf.insured_name,
        "section_reference": sf.policy_reference,
        "underwriter": sf.underwriter,
        "facility_type": ri.facility_type,
        "key_lob": key_lob_str,
        "trifocus": trifocus_str,
        "risk_codes": risk_codes_str,
        "renewal": renewal_str,
        "inception_date": inception_date,
        "expiry_date": expiry_date,
        "line_size": line_size,
        "avg_limit": prem_average_limit,
        "gn_prem": gn_prem,
        "ded_str_1": deductions["ded_str_1"],
        "ded_str_2": deductions["ded_str_2"],
        "ded_str_3": deductions["ded_str_3"],
        "ded_str_4": deductions["ded_str_4"],
        "ded_str_5": deductions["ded_str_5"],
        "ded_str_6": deductions["ded_str_6"],
        "pc": pc_str
    }

    return data
    

def generate_pricing_outputs_dict(hxd, rater):
    cds = hxd.cds
    rs = cds.rating_summary
    pam = rs.pricing_adequacy_metrics

    anti_selection_uncertainty_str = build_anti_selection_uncertainty_str(hxd)

    data = {
        "anti_selection_uncertainty": anti_selection_uncertainty_str,
        "final_pc": pam.pricing_adequacy_actuarial_basis.summary.pc_impact,
        "pc_excl_loads": pam.pricing_adequacy_pre_adj.summary.pc_impact,
        "final_best_estimate": pam.pricing_adequacy_actuarial_basis.summary.best_estimate,
        "best_estimate_excl_loads": pam.pricing_adequacy_pre_adj.summary.best_estimate,
        "final_bpi": pam.pricing_adequacy_actuarial_basis.summary.bpi,
        "bpi_excl_loads": pam.pricing_adequacy_pre_adj.summary.bpi,
        "final_tpi": pam.pricing_adequacy_actuarial_basis.summary.tpi,
        "tpi_excl_loads": pam.pricing_adequacy_pre_adj.summary.tpi,
        "final_roc": pam.pricing_adequacy_actuarial_basis.summary.roc,
        "roc_excl_loads": pam.pricing_adequacy_pre_adj.summary.roc
    }

    return data


def generate_model_gn_ulr_list_dict(hxd, rater):
    projected_gn_ulr_cols = [
        "selected_lob",
        "gn_premium",
        "portfolio_percent",
        "own_exp_gn_ulr",
        "lloyds_gn_ulr",
        "bp_gn_ulr"
    ]
    model_weights_cols = [
        "selected_lob",
        "own_experience",
        "lloyds_proj",
        "bp_proj",
        "model_estimate"
    ]

    model_gn_ulr_df = (
        rater["projected_gn_ulr_df"][projected_gn_ulr_cols]
        .merge(
            rater["model_weights_df"][model_weights_cols],
            how="left",
            on="selected_lob"
        )
    )

    model_gn_ulr_df = model_gn_ulr_df[model_gn_ulr_df["selected_lob"].notna()]

    model_weights_cols.remove("selected_lob")
    model_gn_ulr_df = model_gn_ulr_df[
        projected_gn_ulr_cols + model_weights_cols
    ]

    model_gn_ulr_dict = model_gn_ulr_df.to_dict(orient='records')

    return model_gn_ulr_dict


def generate_selected_ulr_list_dict(hxd, rater):
    model_weights_cols = [
        "selected_lob",
        "model_estimate"
    ]
    cat_loading_cols = [
        "selected_lob",
        "nmp_load_general",
        "nmp_load_weather",
        "climate_change_load"
    ]
    pre_adj_cols = [
        "selected_lob",
        "best_estimate_pre_pc_adj"
    ]
    add_loadings_cols = [
        "selected_lob",
        "anti_selection_charge",
        "uncertainty_charge"
    ]
    act_basis_cols = [
        "selected_lob",
        "pc_impact"
    ]
    final_pricing_cols = [
        "selected_lob",
        "uw_adj",
        "best_estimate_gn"
    ]

    dfs = [
        rater["model_weights_df"][model_weights_cols],
        rater["cat_loadings_df"][cat_loading_cols],
        rater["pre_adj"][pre_adj_cols],
        rater["add_loadings_df"][add_loadings_cols],
        rater["act_basis"][act_basis_cols],
        rater["final_pricing"][final_pricing_cols]
    ]

    merged_df = reduce(
        lambda left, right:
        pd.merge(
            left, 
            right, 
            how="left", 
            on="selected_lob"
        ),
        dfs)

    merged_df = merged_df[merged_df["selected_lob"].notna()]

    merged_df["nmp_load"] = (
        merged_df["nmp_load_general"]
        + merged_df["nmp_load_weather"]
    )

    merged_df = merged_df[[
        "selected_lob",
        "model_estimate",
        "nmp_load",
        "climate_change_load",
        "best_estimate_pre_pc_adj",
        "anti_selection_charge",
        "uncertainty_charge",
        "pc_impact",
        "uw_adj",
        "best_estimate_gn"
    ]]

    selected_ulr_dict = merged_df.to_dict(orient='records')

    return selected_ulr_dict


def generate_tech_adequacy_metrics_list_dict(hxd, rater):
    final_pricing_cols = [
        "selected_lob",
        "best_estimate_gn",
        "bpi",
        "tpi",
        "roc",
    ]
    pre_adj_cols = [
        "selected_lob",
        "best_estimate",
        "bpi",
        "tpi",
        "roc"
    ]
    projected_gn_ulr_cols = [
        "selected_lob",
        "gn_premium"
    ]

    dfs = [
        rater["final_pricing"][final_pricing_cols],
        rater["pre_adj"][pre_adj_cols],
        rater["projected_gn_ulr_df"][projected_gn_ulr_cols],
    ]

    merged_df = reduce(
        lambda left, right:
        pd.merge(
            left, 
            right, 
            how="left", 
            on="selected_lob"
        ),
        dfs)

    merged_df = merged_df[merged_df["selected_lob"].notna()]

    merged_df = merged_df[[
        "selected_lob",
        "gn_premium",
        "best_estimate_gn",
        "bpi_x",
        "tpi_x",
        "roc_x",
        "best_estimate",
        "bpi_y",
        "tpi_y",
        "roc_y",
    ]]

    tech_adequacy_metrics_dict = merged_df.to_dict(orient='records')

    return tech_adequacy_metrics_dict


def generate_model_gn_ulr_summary(hxd):
    model_gn_ulr_path = hxd.cds.rating_summary.model_gn_ulr

    data_dict = {
        "gn_premium_total": model_gn_ulr_path.projected_gn_ulr.summary.gn_premium,
        "portfolio_percent_total": model_gn_ulr_path.projected_gn_ulr.summary.portfolio_percent,
        "own_exp_gn_ulr_total": model_gn_ulr_path.projected_gn_ulr.summary.own_exp_gn_ulr,
        "lloyds_gn_ulr_total": model_gn_ulr_path.projected_gn_ulr.summary.lloyds_gn_ulr,
        "bp_gn_ulr_total": model_gn_ulr_path.projected_gn_ulr.summary.bp_gn_ulr,
        
        "own_experience_total": model_gn_ulr_path.model_weights.summary.own_experience,
        "lloyds_proj_total": model_gn_ulr_path.model_weights.summary.lloyds_proj,
        "bp_proj_total": model_gn_ulr_path.model_weights.summary.bp_proj,
        "model_estimate_total": model_gn_ulr_path.model_weights.summary.model_estimate
    }

    return data_dict


def generate_selected_ulr_summary(hxd):
    rs = hxd.cds.rating_summary
    pam = rs.pricing_adequacy_metrics

    nmp_load_general = rs.cat_loadings.summary.nmp_load_general
    nmp_load_weather = rs.cat_loadings.summary.nmp_load_weather

    nmp_load_total = nmp_load_general + nmp_load_weather

    data_dict = {
        "selected_model_estimate_total": rs.model_gn_ulr.model_weights.summary.model_estimate,
        "nmp_load_total": nmp_load_total,
        "climate_change_total": rs.cat_loadings.summary.climate_change_load,
        "best_estimate_total": pam.pricing_adequacy_pre_adj.summary.best_estimate_pre_pc_adj,
        "anti_selection_total": rs.additional_loadings.summary.anti_selection_charge,
        "uncertainty_total": rs.additional_loadings.summary.uncertainty_charge,
        "pc_impact_total": pam.pricing_adequacy_actuarial_basis.summary.pc_impact,
        "uw_adj_total": pam.pricing_adequacy_final_pricing.summary.uw_adj,
        "best_estimate_gn_total": pam.pricing_adequacy_final_pricing.summary.best_estimate_gn
    }

    return data_dict


def generate_tech_adequacy_metrics_summary(hxd):
    rs = hxd.cds.rating_summary
    pam = rs.pricing_adequacy_metrics

    data_dict = {
        "gn_prem_1": rs.model_gn_ulr.projected_gn_ulr.summary.gn_premium,
        "final_prem": pam.pricing_adequacy_final_pricing.summary.best_estimate_gn,
        "bpi_final": pam.pricing_adequacy_final_pricing.summary.bpi,
        "tpi_final": pam.pricing_adequacy_final_pricing.summary.tpi,
        "roc_final": pam.pricing_adequacy_final_pricing.summary.roc,
        "pre_adj_prem": pam.pricing_adequacy_pre_adj.summary.best_estimate,
        "bpi_pre_adj": pam.pricing_adequacy_pre_adj.summary.bpi,
        "tpi_pre_adj": pam.pricing_adequacy_pre_adj.summary.tpi,
        "roc_pre_adj": pam.pricing_adequacy_pre_adj.summary.roc
    }

    return data_dict


def hide_empty_rows(wb, start_row, end_row, tbl):
    worksheet = wb['Note Tables']
    num_row = len(tbl)
    start_hidden_row = num_row + start_row
    if start_hidden_row < end_row:      
        worksheet.row_dimensions.group(
            start = start_hidden_row,
            end = end_row,
            hidden= True
        )


def write_preserve_style(sheet, row, col, value):
    cell = sheet.cell(row=row, column=col)

    font = copy(cell.font)
    fill = copy(cell.fill)
    border = copy(cell.border)
    alignment = copy(cell.alignment)
    number_format = copy(cell.number_format)
    protection = copy(cell.protection)

    cell.value = value

    cell.font = font
    cell.fill = fill
    cell.border = border
    cell.alignment = alignment
    cell.number_format = number_format
    cell.protection = protection

