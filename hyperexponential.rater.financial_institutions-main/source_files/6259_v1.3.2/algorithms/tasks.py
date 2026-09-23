import os
import re
from io import BytesIO

import hx
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from mailmerge import MailMerge
from algorithms.task_rate_change import rarc, expiring_policy_fetch
from algorithms.climate_documents.generate_climate_doc import generate_climate_doc
from algorithms.climate_documents.import_climate_docs import FeatherLoader
from datetime import date
from tempfile import NamedTemporaryFile

def is_hx_list_like(obj):
    """Return True if obj behaves like an HX List."""

    HX_LIST_METHODS = [
        '__contains__', '__delitem__', '__getitem__', '__iadd__', '__iter__', '__len__', '__reversed__', 'append',
        'count', 'extend', 'index', 'insert', 'remove'
    ]

    return all(hasattr(obj, attr) for attr in HX_LIST_METHODS)


def get_from_path(obj, path, excel_name):
    """
    Resolve a value from `obj` using a CDS-style path.
    Injects list indexing using LayerN pattern from the Excel named range.
    """
    parts = path.strip("/").split("/")

    layer_match = re.search(r"Layer(\d+)", excel_name)
    layer_index = int(layer_match.group(1)) - 1 if layer_match else None

    for attr in parts:
        if hasattr(obj, attr):
            next_obj = getattr(obj, attr)

            if is_hx_list_like(next_obj) and layer_index is not None:
                if layer_index < len(next_obj):
                    obj = next_obj[layer_index]
                else:
                    return None
                continue

            obj = next_obj
        else:
            return None

    return obj


@hx.task
def export_to_excel(hxd, progress):
    """
    Export input HXD values into the Excel template using named ranges.
    Supports resolving HX list entries based on naming convention LayerN,
    and toggles data source (values vs. percent) for region-based metrics
    based on user selection in HXD.
    """

    # --- File paths ---

    base_path = os.path.dirname(__file__)
    wb_path = os.path.join(base_path, "excel_export/template_rater.xlsx")
    mappings_path = os.path.join(base_path, "excel_export/mappings.csv")

    # --- Load workbook and mappings ---

    workbook = load_workbook(wb_path, keep_vba=False)
    mappings = pd.read_csv(mappings_path)

    # --- Region metric toggling (values vs percent) ---

    regions = [
        "africa", "caribbean", "arab_states", "asia", "oceania", "europe",
        "former_soviet_republics", "usa", "canada", "south_latin_america", "row"
    ]
    metrics = ["employees", "assets", "revenues"]

    for region in regions:
        for metric in metrics:
            base_path = f"cds/exposure/granular/regions/regions_list/{region}"
            value_path = f"{base_path}/values/{metric}"
            percent_path = f"{base_path}/percent/{metric}"

            is_value_selected = getattr(hxd.non_cds.exposure_details.regions.is_value, metric)
            is_percent_selected = getattr(hxd.non_cds.exposure_details.regions.is_percent, metric)

            if is_value_selected:
                mappings.loc[mappings["CDS Path"] == percent_path, "Add to excel"] = False
                mappings.loc[mappings["CDS Path"] == value_path, "Add to excel"] = True
            elif is_percent_selected:
                mappings.loc[mappings["CDS Path"] == value_path, "Add to excel"] = False
                mappings.loc[mappings["CDS Path"] == percent_path, "Add to excel"] = True

    # --- Risk assessment toggling based on `shownBy` conditions ---

    risk_assessment_fields = {
        "policy_wording": "any_any",
        "claims_history_cpi": "any_crime_pi",
        "claims_history_do": "any_do",
        "risk_management": "any_any",
        "strength_of_financial": "any_any",
        "technological_infrastructure": "any_any",
        "quality_of_control": "any_crime",
        "agents_as_employees": "notins_crime_pi",
        "regulatory_risk": "any_pi_do",
        "quality_of_claims_handling": "ins_pi",
        "product_complexity": "any_pi",
        "quality_of_bcp": "fin_pi",
        "market_regulator": "ban_fin_pi",
        "extent_of_leveraged_gearing": "inv_pi",
        "quality_of_performance_non_pevc": "inv_pi",
        "redemption_gates": "inv_pi",
        "valuation_for_pevc": "inv_PE_VC_RE_pi",
        "loan_covenant": "inv_PE_VC_RE_pi",
        "dando_portfolio_companies": "inv_PE_VC_do",
        "data_centre": "fin_any",
        "tech_outsourcing": "fin_any",
    }

    for field, condition in risk_assessment_fields.items():
        condition = getattr(hxd.non_cds.risk_assesment,condition)
    
        if condition is False:
            mappings.loc[mappings["CDS Path"] == f"cds/modifiers/comment/{field}", "Add to excel"] = False
            mappings.loc[mappings["CDS Path"] == f"cds/modifiers/risk_category/{field}", "Add to excel"] = False


    # --- Write values to Excel ---

    for _, row in mappings.iterrows():
        try:
            mode = str(row.get("Mode", "")).lower()

            # Skip output rows
            if str(row.get("Add to excel", "true")).lower() != "true":
                continue

            excel_name = row["Template Excel rater named range"]
            cds_path = row["CDS Path"]

            # Special case for policy option ID
            if excel_name == "DB.hxOptionID":
                value = hx.meta.policy_option_id
            else:
                # If override mode, append '.selected' to the last path element
                if mode == "override":
                    cds_path = cds_path.rstrip("/") + "/selected"

                value = get_from_path(hxd, cds_path, excel_name)


            if value is None:
                continue

            if "excess_str" in cds_path:
                try:
                    value = float(str(value).replace(",", ""))
                except (ValueError, TypeError):
                    value = 0.0

            named_range = workbook.defined_names[excel_name]

            for sheet_name, coord in named_range.destinations:
                sheet = workbook[sheet_name]
                sheet[coord].value = value

        except Exception as e:
            print(e)

    # --- Protect all sheets ---
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(
            min_row=sheet.min_row,
            max_row=sheet.max_row,
            min_col=sheet.min_column,
            max_col=sheet.max_column
        ):
            for cell in row:
                if cell.protection is not None and not cell.protection.locked:
                    cell.protection = cell.protection.copy(locked=True)

        sheet.protection.sheet = True
        sheet.protection.enable()
        sheet.protection.set_password("Safe")


    # --- Save workbook ---
    with hxd.output_file.open("b") as f:
        workbook.save(f)


@hx.task
def num_layers_task(hxd, progress):
    current_num_layers = len(hxd.cds.layers)
    new_num_layers = hxd.non_cds.cover_details.num_layers

    if current_num_layers < new_num_layers:
        # Add new rows
        hxd.cds.layers.extend([{} for _ in range(current_num_layers, new_num_layers)])
    elif current_num_layers > new_num_layers:
        # Delete rows
        for i in range(current_num_layers - 1, new_num_layers - 1, -1):
            del hxd.cds.layers[i]


@hx.task
def sync_expiring_ids(hxd, progress):
    # Note this task is to be run following each migrated policy to sync the expiring ids
    hxd.model_state.expiring_policy_option_id = hx.meta.expiring_policy_option_id



@hx.task
def start_renewal_task(hxd, progress):
    # The following statement checks that the expiring information has been imported for renewals. Some teams might want to start
    # from a blank rater each time, in which case update the below
    ms = hxd.model_state
    sf = hxd.cds.standard_fields
    if not hxd.cds.standard_fields.insured_name:
        ms.landing_page_info = "❗**FAILED**: Click 'Undo' then 'Import Expiring Policy Data' in the top right corner.❗"
    else:
        ms.pressed_start_renewal_task = True
        ms.expiring_policy_option_id = hx.meta.expiring_policy_option_id
        sf.is_renewal = True
        # Add tasks which must be done before starting a policy here >>


# Importing expiring policy for rate change
@hx.task
def expiring_policy_fetch_task(hxd, progress):
    expiring_policy_fetch(hxd, progress)


@hx.task
def rarc_task(hxd, progress):
    rarc(hxd, progress)


@hx.task
def save_quote_details_to_pas_reference(hxd, progress):
    pas_reference = hxd.cds.risk_info.quote_details_search
    hx.meta.pas_references.clear()
    hx.meta.pas_references.append(pas_reference)


@hx.task
def generate_tags(hxd, progress):
    # Tags for Underwriter, Broker, CountryOfDomicile, Industry, SubIndustry
    tag_fields = {
        'Underwriter': hxd.cds.standard_fields.underwriter,
        'Broker': hxd.cds.standard_fields.broker,
        'Country': hxd.cds.standard_fields.insured_country,
        'Industry': hxd.cds.key_industry.code_name,
        'Sub Industry': hxd.cds.rating_factors.risk_info.sub_industry
    }

    new_tags = []
    tag_fields_error = []

    # Validate and generate tags
    for label, value in tag_fields.items():
        if value:
            new_tags.append(value.replace(" ", ""))
        else:
            tag_fields_error.append(label)

    # Set policy tags or error message
    if tag_fields_error:
        msg_string = 'Error saving policy tags - Please select: ' + ', '.join(tag_fields_error)
        hx.errors.validation(msg_string)
        hxd.cds.is_policy_tag_error = True
        hxd.cds.policy_tag_msg = msg_string
    else:
        hx.meta.policy_tags = new_tags
        hxd.cds.is_policy_tag_error = False
        hxd.cds.policy_tag_msg = "Success!"


def _normalise_field_name(field_name: str) -> str:
    return "".join(ch for ch in field_name.lower() if ch.isalnum())
 

@hx.task
def generate_climate_doc_task(hxd, progress):
    country = hxd.cds.climate_document_country
    endpoint = (
                    "https://beazley.sharepoint.com/sites/ClimateRisk/"
                    "Shared Documents/"
                    "4. Underwriting and Pricing/"
                    "7. Specialty ESG Group/"
                    "Climate Litigation Spotlight/"
                    "Spotlight Drafts"
                )

    loader = FeatherLoader(endpoint=endpoint)
    document_paths = {
        "Australia": "/Australia/Litigation Spotlight AUS_FI.docx",
        "France": "/France/Litigation Spotlight France_FI.docx",
        "Germany": "/Germany/Litigation Spotlight Germany_FI.docx",
        "The Netherlands": "/The Netherlands/Litigation Spotlight Nlds_FI.docx",
        "United Kingdom": "/United Kingdom/Litigation Spotlight UK_FI.docx",
    }

    try:
        relative_path = document_paths[country]
    except KeyError as exc:
        hx.errors.fatal(f"No template configured for country '{country}'")
        raise exc

    doc = loader.load_word_document(relative_path)

    with NamedTemporaryFile(suffix=".docx") as tmp:
        doc.save(tmp.name)
        document = MailMerge(tmp.name)

        document.merge(
            insured_name=str(hxd.cds.standard_fields.insured_name or ""),
            date=date.today().strftime("%d %B %Y")
        )

        with hxd.cds.climate_document.open("b") as f:
            document.write(f)

    progress.update(1.0)
