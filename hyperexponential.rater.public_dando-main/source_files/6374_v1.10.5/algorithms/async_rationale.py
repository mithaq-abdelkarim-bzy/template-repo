import hx
import pandas as pd
import numpy as np
import math as math
import algorithms.rate_utilities as utils
import algorithms.rate_constants as const
from algorithms import parameter_tables_schema as params
from docx import Document

def process_comment(hxd, progress):
   
    comments = {
        "Underwriter Rationale" : hxd.cds.standard_fields.uw_rationale,
        "General Comments" : hxd.cds.comment,
        "Claims History" : hxd.cds.claims_history,
        "T&C's" : hxd.cds.t_and_c_comment,
        "Management and Corporate Governance" : hxd.cds.modifiers.management_corp_gov.comment,
        "Business / Financial Model Factors" : hxd.cds.modifiers.business_financial_model_factors.comment,
        "Significant Event Factors" : hxd.cds.modifiers.significant_event_factors.comment,
        "Stock Market Factors" : hxd.cds.modifiers.stock_market_factors.comment,
        "Derivative" : hxd.cds.modifiers.derivative.comment,
        "Regulatory" : hxd.cds.modifiers.regulatory.comment,
        "M&A" : hxd.cds.modifiers.ma.comment,
        "Underwriter Override Comment" : hxd.cds.key_industry.sector_uw_override_comment,
        "Underwriter Override Comment 2" : hxd.cds.key_industry.sector_uw_override_comment_2,
        "Underwriter Override Comment 3" : hxd.cds.key_industry.sector_uw_override_comment_3,
        "Market Cap Comment" : hxd.cds.exposure.aggregate.revised_market_cap_comment,
        "IPO Date Comment" : hxd.cds.exposure.aggregate.ipo_date_comment
    }

    comment = comments[hxd.cds.comment_tool.input_comment]

    # Create a new Document
    doc = Document()

    # Add a title to the document
    doc.add_heading(hxd.cds.comment_tool.input_comment, 0)

    # Add a paragraph
    doc.add_paragraph(comment)

    # Save to the hxd
    with hxd.cds.comment_tool.document.open("b") as f:
        doc.save(f)

    # Reduce the comment size
    if comment is not None:
        if len(comment) >= 20000:
            comment_adj = comment[:10000]
            setattr(hxd.cds, "comment", comment_adj)

    pass
